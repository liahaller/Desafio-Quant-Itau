"""Reestiliza o deck da final no estilo do template T6 (Final/Reestilizar).

Lê `Final/FINAL.pptx` e grava `Final/FINAL_T6.pptx`. Mexe só em estilo —
fundo, fontes, cores, cantos, efeitos, cabeçalho e rodapé — e preserva todo
o conteúdo (textos, posições, imagens, transições Morph):

  fundo      radial navy + halos + número fantasma  →  grafite chapado (15181D)
  cabeçalho  "KAIROS · seção" + nº + filete + título abaixo
             →  "Seção | Título" numa linha, régua, marca KAIRÓS à direita
  rodapé     "ITAÚ QUANT AI CHALLENGE · FINAL 2026 / BLACK-LITTERMAN × POLYMARKET"
             →  barra de navegação (seções do deck, ativa em destaque) + nº da página
  cards      roundRect em degradê com glow  →  retângulo chapado, borda 1 pt cinza
  fontes     Bahnschrift → Segoe UI; Consolas só nos kickers → Segoe UI Semibold
             (fórmulas, nomes de arquivo e código continuam em Consolas)
  cores      texto/linhas navy → cinzas do T6; azul neon → azul Polymarket escurecido

Uso:
    python scripts/final_t6_pptx.py            # gera Final/FINAL_T6.pptx
    python scripts/final_t6_pptx.py --render   # e exporta PNG (scratchpad) via PowerPoint
"""

import copy
import os
import re
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
ENTRADA = RAIZ / "Final" / "FINAL.pptx"
SAIDA = RAIZ / "Final" / "FINAL_T6.pptx"

# ------------------------------------------------------------------ paleta T6
BG = "15181D"
PANEL = "1E2229"
RULE = "3A404A"
TEXT = "FFFFFF"
BODY = "D5D9E0"
MUTED = "9AA3AF"
FOOT = "6B7280"
GOLD = "FFB531"
BLUE_FILL = "2A4BC4"
BLUE_TEXT = "6C8CE0"
BLUE_LINE = "4F6FD0"

COR_TEXTO = {"F5F7FB": TEXT, "BAC6DA": BODY, "8395B5": MUTED, "58698A": FOOT,
             "8FA2C2": MUTED, "4F8EF7": BLUE_TEXT}
COR_LINHA = {"2A3F63": RULE, "35507A": RULE, "4F8EF7": BLUE_LINE, "8395B5": MUTED}
COR_FILL = {"13213A": PANEL, "0B1426": PANEL, "14233E": PANEL, "182946": PANEL,
            "0F182B": PANEL, "04070E": BG, "4F8EF7": BLUE_FILL, "2A3F63": RULE,
            "22345A": RULE}

# fonte de origem → (fonte T6, fator de tamanho). Segoe UI é mais larga que a
# Bahnschrift (e muito mais que a condensada): o fator segura as quebras.
FONTES = {
    "Bahnschrift SemiBold Condensed": ("Segoe UI Semibold", 0.80),
    "Bahnschrift SemiBold": ("Segoe UI Semibold", 0.95),
    "Bahnschrift SemiLight": ("Segoe UI", 0.95),
    "Bahnschrift Light": ("Segoe UI", 0.95),
    "Bahnschrift": ("Segoe UI", 0.95),
}
KICKER = ("Segoe UI Semibold", 1.0)        # Consolas em caixa alta

F_HDR = "Segoe UI Semibold"
F_NAV = "Segoe UI"

# Seções da barra de navegação (vocabulário do próprio deck) e o slide em que
# cada uma começa. ponytail: lista fixa — mudar aqui se o deck for reordenado.
NAV = [("Hipótese", 1), ("Polymarket", 3), ("Estratégia", 5), ("Backtest", 14),
       ("IA", 16), ("Conclusão", 31)]

L, R = 0.80, 12.53                        # margens dominantes do deck
A = nsdecls("a")


# --------------------------------------------------------------- utilidades
def pol(emu):
    return (emu or 0) / 914400


def texto(sh):
    return sh.text_frame.text if sh.has_text_frame else ""


def fonte_run(r):
    lat = r.find(qn("a:rPr") + "/" + qn("a:latin"))
    return lat.get("typeface") if lat is not None else None


def tamanho_run(r):
    rpr = r.find(qn("a:rPr"))
    return int(rpr.get("sz")) if rpr is not None and rpr.get("sz") else None


def primeira_fonte(sh):
    for r in sh._element.iter(qn("a:r")):
        return fonte_run(r), tamanho_run(r)
    return None, None


def remover(sh):
    sh._element.getparent().remove(sh._element)


def eh_linha_fina(sh):
    """Filete horizontal: conector `line` ou retângulo de altura ~0."""
    return pol(sh.width) > 6 and pol(sh.height) < 0.03


# ------------------------------------------------------------ classificação
def classificar(sh):
    """Devolve o papel do shape no chrome do deck atual, ou None."""
    nome = sh.name.lower()
    y, h = pol(sh.top), pol(sh.height)
    t = texto(sh).strip()
    f, sz = primeira_fonte(sh)
    if "glow" in nome or "halo" in nome or "ghost" in nome:
        return "decor"
    if sz and sz >= 14000:                 # número fantasma (150 pt); o KAIROS da capa tem 120
        return "decor"
    if nome.endswith("tick") or (y < 0.6 and pol(sh.width) < 0.08 and h < 0.25 and not t):
        return "decor"
    if y < 0.6 and h < 0.3:
        if t == "KAIROS":
            return "hdr_marca"
        if t.startswith("KAIROS"):
            return "hdr_secao"
        if re.fullmatch(r"\d{1,2}", t):
            return "hdr_num"
        if f == "Consolas" and t:
            return "hdr_secao_solta"
    if 0.7 <= y <= 0.85 and eh_linha_fina(sh):
        return "hdr_regua"
    if 0.85 <= y <= 1.0 and sz and 1800 <= sz <= 2200 and t:
        return "titulo"
    if y > 7.1 and eh_linha_fina(sh):
        return "foot_regua"
    if y > 7.1 and h < 0.3 and f == "Consolas":
        return "foot_txt"
    return None


# ------------------------------------------------------------- reestilizar
def fundo(slide):
    cSld = slide._element.find(qn("p:cSld"))
    bg = cSld.find(qn("p:bg"))
    if bg is not None:
        cSld.remove(bg)
    cSld.insert(0, parse_xml(
        f'<p:bg {nsdecls("p")} {A}><p:bgPr><a:solidFill><a:srgbClr val="{BG}"/></a:solidFill>'
        f"<a:effectLst/></p:bgPr></p:bg>"))


def recolorir_spPr(sp_pr):
    """Preenchimento, linha e efeitos de um shape."""
    geom = sp_pr.find(qn("a:prstGeom"))
    if geom is not None and geom.get("prst") == "roundRect":
        geom.set("prst", "rect")
        av = geom.find(qn("a:avLst"))
        if av is not None:
            av.clear()
    eff = sp_pr.find(qn("a:effectLst"))
    if eff is not None:
        eff.clear()
    grad = sp_pr.find(qn("a:gradFill"))
    if grad is not None:
        cores = {c.get("val") for c in grad.iter(qn("a:srgbClr"))}
        if cores <= {"182946", "0F182B"}:               # card em degradê → chapado
            sp_pr.replace(grad, parse_xml(
                f'<a:solidFill {A}><a:srgbClr val="{PANEL}"/></a:solidFill>'))
        else:
            for c in grad.iter(qn("a:srgbClr")):
                if c.get("val") in ("4F8EF7",):
                    c.set("val", BLUE_FILL)
    sol = sp_pr.find(qn("a:solidFill") + "/" + qn("a:srgbClr"))
    if sol is not None and sol.get("val") in COR_FILL:
        sol.set("val", COR_FILL[sol.get("val")])
    ln = sp_pr.find(qn("a:ln"))
    if ln is not None:
        for c in ln.iter(qn("a:srgbClr")):
            if c.get("val") in COR_LINHA:
                c.set("val", COR_LINHA[c.get("val")])
        lg = ln.find(qn("a:gradFill"))
        if lg is not None:                                # borda em degradê → sólida
            ln.replace(lg, parse_xml(f'<a:solidFill {A}><a:srgbClr val="{RULE}"/></a:solidFill>'))


def refontar(el):
    """Troca fontes, escala tamanhos e recolore os runs de um elemento."""
    for rpr in list(el.iter(qn("a:rPr"))) + list(el.iter(qn("a:endParaRPr"))):
        lat = rpr.find(qn("a:latin"))
        sz = rpr.get("sz")
        if lat is not None:
            f = lat.get("typeface")
            if f in FONTES:
                nova, fator = FONTES[f]
            elif f == "Consolas":
                r = rpr.getparent()
                t = r.findtext(qn("a:t")) or ""
                if t and t == t.upper() and re.search(r"[A-ZÀ-Ú]", t):
                    nova, fator = KICKER
                else:
                    nova, fator = None, 1.0
            else:
                nova, fator = None, 1.0
            if nova:
                lat.set("typeface", nova)
                for tag in ("a:ea", "a:cs"):
                    x = rpr.find(qn(tag))
                    if x is not None:
                        x.set("typeface", nova)
                if sz and fator != 1.0:
                    rpr.set("sz", str(max(600, int(round(int(sz) * fator / 50.0)) * 50)))
        c = rpr.find(qn("a:solidFill") + "/" + qn("a:srgbClr"))
        if c is not None and c.get("val") in COR_TEXTO:
            c.set("val", COR_TEXTO[c.get("val")])


def caixa(slide, nome, x, y, w, h, anchor="ctr"):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = nome
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf._txBody.find(qn("a:bodyPr")).set("anchor", anchor)
    return sh


def run(p, t, fonte, pt, cor, negrito=False, spc=None):
    r = p.add_run()
    r.text = t
    r.font.name, r.font.size, r.font.bold = fonte, Pt(pt), negrito
    r.font.color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor.from_string(cor)
    if spc:
        r._r.get_or_add_rPr().set("spc", str(spc))
    return r


def linha(slide, nome, x, y, w, cor):
    sh = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(0.01))  # 1 = rect
    sh.name = nome
    sh.fill.solid()
    sh.fill.fore_color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor.from_string(cor)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def margens(slide):
    """L/R do slide = caixa do conteúdo que sobrou (chrome já removido), com folga."""
    xs = [pol(sh.left) for sh in slide.shapes if sh.left is not None and 0.3 <= pol(sh.left) <= 1.2]
    xr = [pol(sh.left + sh.width) for sh in slide.shapes
          if sh.left is not None and 12.0 <= pol(sh.left + sh.width) <= 13.0]
    return (min(xs) if xs else L), (max(xr) if xr else R)


def cabecalho(slide, secao, titulo_runs, numero, L=L, R=R):
    """'Seção | Título' + régua + marca. `titulo_runs` = lista de (texto, cor)."""
    from pptx.enum.text import PP_ALIGN
    n = len(secao) + sum(len(t) for t, _ in titulo_runs)
    pt = 20 if n <= 62 else 18 if n <= 76 else 16 if n <= 92 else 14
    sh = caixa(slide, "!!t6_hdr", L, 0.38, R - L - 1.6, 0.5)
    p = sh.text_frame.paragraphs[0]
    run(p, secao, F_HDR, pt, TEXT, negrito=True)
    if titulo_runs:
        run(p, "  |  ", F_HDR, pt, MUTED)
        for t, cor in titulo_runs:
            run(p, t, F_HDR, pt, cor)
    m = caixa(slide, "!!t6_marca", R - 1.5, 0.38, 1.5, 0.5)
    pm = m.text_frame.paragraphs[0]
    pm.alignment = PP_ALIGN.RIGHT
    run(pm, "KAIRÓS", F_HDR, 11, GOLD, spc=300)
    linha(slide, "!!t6_regua", L, 0.95, R - L, RULE)


def rodape(slide, idx, numero, L=L, R=R):
    from pptx.enum.text import PP_ALIGN
    linha(slide, "!!t6_nav_regua", L, 7.14, R - L, RULE)   # o conteúdo do deck desce até ~7,05
    ativa = max((s for s, ini in NAV if idx >= ini), key=lambda s: dict(NAV)[s])
    n = len(NAV)
    wtot = R - L - 0.8
    for i, (s, _) in enumerate(NAV):
        xi = L + i * wtot / n
        sh = caixa(slide, f"!!t6_nav{i}", xi, 7.18, wtot / n, 0.26)
        p = sh.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        eh = s == ativa
        run(p, s, F_HDR if eh else F_NAV, 8.5, TEXT if eh else MUTED)
        if eh:
            sub = linha(slide, "!!t6_nav_ativa", xi + wtot / n / 2 - 0.45, 7.45, 0.9, GOLD)
            sub.height = Inches(0.03)
    num = caixa(slide, "!!t6_num", R - 0.65, 7.18, 0.65, 0.26)
    pn = num.text_frame.paragraphs[0]
    pn.alignment = PP_ALIGN.RIGHT
    run(pn, f"{idx:02d}", F_NAV, 9, MUTED)   # posição real; os nºs antigos estavam defasados


def reestilizar(slide, idx, avisos):
    fundo(slide)
    secao, numero, titulo_runs = "", None, []
    for sh in list(slide.shapes):
        papel = classificar(sh)
        if papel is None:
            continue
        if papel == "hdr_secao":
            secao = texto(sh).strip()[len("KAIROS"):].strip()
        elif papel == "hdr_secao_solta":
            secao = texto(sh).strip()
        elif papel == "hdr_num":
            numero = texto(sh).strip()
        elif papel == "titulo":
            for r in sh._element.iter(qn("a:r")):
                t = r.findtext(qn("a:t")) or ""
                c = r.find(qn("a:rPr") + "/" + qn("a:solidFill") + "/" + qn("a:srgbClr"))
                cor = c.get("val") if c is not None else "F5F7FB"
                titulo_runs.append((t, COR_TEXTO.get(cor, cor)))
        remover(sh)

    for sh in slide.shapes:
        el = sh._element
        sp_pr = el.find(qn("p:spPr"))
        if sp_pr is not None:
            recolorir_spPr(sp_pr)
        refontar(el)
        if sh.shape_type == 13:                           # imagem: só avisa se tiver fundo opaco
            from PIL import Image
            import io
            im = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA")
            px = im.getpixel((0, 0))
            if px[3] == 255 and max(px[:3]) > 40:
                avisos.append(f"slide {idx}: imagem '{sh.name}' com fundo opaco {px[:3]}")

    if titulo_runs and "".join(t for t, _ in titulo_runs).strip().lower() in secao.lower():
        titulo_runs = []                    # 'A hipótese | Hipótese' → só a seção
    l, r = margens(slide)
    cabecalho(slide, secao, titulo_runs, numero, l, r)
    rodape(slide, idx, numero, l, r)


def render(arquivo, pasta):
    pasta.mkdir(parents=True, exist_ok=True)
    ps = (f"$app = New-Object -ComObject PowerPoint.Application; "
          f"$p = $app.Presentations.Open('{arquivo}', $true, $false, $false); "
          f"for ($i=1; $i -le $p.Slides.Count; $i++) {{ $p.Slides.Item($i).Export("
          f"('{str(pasta).replace(chr(92), '/')}/s{{0:d2}}.png' -f $i).Replace('/','\\'), 'PNG', 1280, 720) }}; "
          f"$p.Close(); $app.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)


if __name__ == "__main__":
    prs = Presentation(ENTRADA)
    avisos = []
    for i, s in enumerate(prs.slides, 1):
        reestilizar(s, i, avisos)
    prs.save(SAIDA)
    print("ok", SAIDA.relative_to(RAIZ), f"({len(prs.slides)} slides)")
    for a in avisos:
        print("aviso:", a)
    if "--render" in sys.argv:
        pasta = Path(os.environ.get("RENDER_DIR", RAIZ / "Final" / "render_t6"))
        render(SAIDA, pasta)
        print("PNG em", pasta)

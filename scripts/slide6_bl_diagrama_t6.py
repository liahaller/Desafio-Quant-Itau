"""Refaz o slide 6 do `Final/FINAL_T6.pptx` como o diagrama de Idzorek (2005).

Replica a figura clássica do Black-Litterman (caixas de entrada → π →
distribuição a priori · views + Ω → distribuição das views → distribuição
combinada, com as sininhas) na notação do projeto (δ, τ = 1/504, 9 ETFs,
k views) e na pele T6. Cabeçalho, marca, régua e barra de navegação do
slide (`!!t6_*`) ficam; todo o resto do slide 6 é apagado e redesenhado.

Uso:
    python scripts/slide6_bl_diagrama_t6.py            # grava no próprio FINAL_T6.pptx
    python scripts/slide6_bl_diagrama_t6.py --render   # e exporta o slide 6 em PNG
"""

import math
import os
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
ARQUIVO = RAIZ / "Final" / "FINAL_T6.pptx"
SLIDE = 6

# paleta T6 (mesma de final_t6_pptx.py)
PANEL, RULE = "1E2229", "3A404A"
TEXT, BODY, MUTED, FOOT, GOLD = "FFFFFF", "D5D9E0", "9AA3AF", "6B7280", "FFB531"
F_SEMI, F_REG = "Segoe UI Semibold", "Segoe UI"

L, R = 0.90, 12.43            # margens do slide 6 (régua do cabeçalho)


# ---------------------------------------------------------------- primitivas
def rgb(c):
    return RGBColor.from_string(c)


def caixa_texto(slide, nome, x, y, w, h, anchor="ctr"):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = nome
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf._txBody.find(qn("a:bodyPr")).set("anchor", anchor)
    return sh


def run(p, t, pt, cor=TEXT, fonte=F_SEMI, sub=False, sup=False):
    r = p.add_run()
    r.text = t
    r.font.name, r.font.size = fonte, Pt(pt)
    r.font.color.rgb = rgb(cor)
    if sub or sup:
        r._r.get_or_add_rPr().set("baseline", "30000" if sup else "-25000")
    return r


def paragrafo(sh, alinhar=PP_ALIGN.CENTER, primeiro=True):
    tf = sh.text_frame
    p = tf.paragraphs[0] if primeiro else tf.add_paragraph()
    p.alignment = alinhar
    return p


def retangulo(slide, nome, x, y, w, h):
    sh = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))   # 1 = rect
    sh.name = nome
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(PANEL)
    sh.line.color.rgb = rgb(RULE)
    sh.line.width = Pt(1)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def seta(slide, nome, x, y0, y1):
    """Seta vertical fina, ponta em baixo."""
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y0), Inches(x), Inches(y1))
    c.name = nome
    c.line.color.rgb = rgb(MUTED)
    c.line.width = Pt(1)
    ln = c._element.find(qn("p:spPr")).find(qn("a:ln"))
    ln.append(parse_xml(f'<a:tailEnd {nsdecls("a")} type="triangle" w="med" len="med"/>'))
    return c


def sininha(slide, nome, x, y, w, h, cor, largura_pt=1.5, sigma=0.16):
    """Curva normal (só a linha) dentro da caixa (x, y, w, h) + linha de base."""
    n = 80
    pts = []
    for i in range(n + 1):
        u = i / n
        z = (u - 0.5) / sigma
        pts.append((Inches(x + u * w), Inches(y + h * (1 - math.exp(-z * z / 2)))))
    ff = slide.shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
    ff.add_line_segments(pts[1:], close=False)
    sh = ff.convert_to_shape()
    sh.name = nome
    sh.fill.background()
    sh.line.color.rgb = rgb(cor)
    sh.line.width = Pt(largura_pt)
    base = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x - 0.1), Inches(y + h + 0.06),
                                      Inches(x + w + 0.1), Inches(y + h + 0.06))
    base.name = nome + "_base"
    base.line.color.rgb = rgb(RULE)
    base.line.width = Pt(0.75)
    return sh


# ------------------------------------------------------------------- blocos
def caixa_entrada(slide, nome, x, y, w, h, titulo, simbolo, dim):
    """Caixa de entrada: título · símbolo entre parênteses grandes · dimensão."""
    retangulo(slide, nome, x, y, w, h)
    t = caixa_texto(slide, nome + "_t", x + 0.08, y + 0.07, w - 0.16, 0.22, "t")
    run(paragrafo(t), titulo, 9.5, BODY, F_REG)
    s = caixa_texto(slide, nome + "_s", x + 0.08, y + 0.30, w - 0.16, 0.34)
    p = paragrafo(s)
    for txt, pt, cor, sub in simbolo:
        run(p, txt, pt, cor, sub=sub)
    d = caixa_texto(slide, nome + "_d", x + 0.08, y + h - 0.24, w - 0.16, 0.18, "b")
    run(paragrafo(d), dim, 8, MUTED, F_REG)


def caixa_distribuicao(slide, nome, x, y, w, h, titulo, formula, cor_curva, rodape=None, sigma=0.16, largura=1.5):
    retangulo(slide, nome, x, y, w, h)
    t = caixa_texto(slide, nome + "_t", x + 0.15, y + 0.10, w - 0.30, 0.22, "t")
    run(paragrafo(t), titulo, 10, BODY, F_REG)
    alt = h - 0.40 - 0.42 - (0.22 if rodape else 0.05)
    sininha(slide, nome + "_curva", x + 0.55, y + 0.40, w - 1.10, alt - 0.10, cor_curva, largura, sigma)
    f = caixa_texto(slide, nome + "_f", x + 0.15, y + 0.40 + alt, w - 0.30, 0.34)
    p = paragrafo(f)
    for txt, pt, cor, sub in formula:
        run(p, txt, pt, cor, sub=sub)
    if rodape:
        r = caixa_texto(slide, nome + "_r", x + 0.15, y + h - 0.26, w - 0.30, 0.18, "b")
        run(paragrafo(r), rodape, 8, MUTED, F_REG)


# ---------------------------------------------------------------------- slide
PAR = 24                       # tamanho dos parênteses "de matriz"
SYM = 17


def par(simb, cor=TEXT):
    """( símbolo ) com parênteses maiores, como na figura de Idzorek."""
    return [("( ", PAR, MUTED, False), (simb, SYM, cor, False), (" )", PAR, MUTED, False)]


def refazer(slide):
    for sh in list(slide.shapes):
        if not sh.name.startswith("!!t6_"):
            sh._element.getparent().remove(sh._element)

    # título ao lado da seção (o cabeçalho do T6 é "Seção | Título")
    hdr = next(sh for sh in slide.shapes if sh.name == "!!t6_hdr")
    p = hdr.text_frame.paragraphs[0]
    pt = p.runs[0].font.size.pt
    run(p, "  |  ", pt, MUTED)
    run(p, "Do equilíbrio às views", pt, TEXT)

    # --- linha 1: entradas
    y1, h1, w1 = 1.22, 0.84, 2.05
    xs_esq = [0.90, 3.10, 5.30]
    xs_dir = [8.33, 10.48]
    caixa_entrada(slide, "s6_delta", xs_esq[0], y1, w1, h1, "Aversão ao risco",
                  [("δ = (E[r] − r", 13, TEXT, False), ("f", 9, TEXT, True), (") / σ²", 13, TEXT, False)],
                  "escalar  ·  δ = 3,0")
    caixa_entrada(slide, "s6_sigma", xs_esq[1], y1, w1, h1, "Matriz de covariância",
                  par("Σ"), "9 × 9  ·  amostral")
    caixa_entrada(slide, "s6_wmkt", xs_esq[2], y1, w1, h1, "Pesos de capitalização",
                  [("( ", PAR, MUTED, False), ("w", SYM, TEXT, False), ("mkt", 11, TEXT, True), (" )", PAR, MUTED, False)],
                  "9 × 1  ·  soma 1")
    caixa_entrada(slide, "s6_views", xs_dir[0], y1, 1.95, h1, "Views",
                  [("( ", PAR, MUTED, False), ("P", SYM, GOLD, False), (", ", SYM, TEXT, False),
                   ("Q", SYM, GOLD, False), (" )", PAR, MUTED, False)],
                  "P: k × 9  ·  Q: k × 1")
    caixa_entrada(slide, "s6_omega", xs_dir[1], y1, 1.95, h1, "Incerteza das views",
                  par("Ω", GOLD), "k × k  ·  diagonal")

    # --- linha 2: retorno de equilíbrio implícito
    y2, h2 = 2.38, 0.66
    x2, w2 = (xs_esq[0] + xs_esq[-1] + w1) / 2 - 2.90, 5.80
    for x in xs_esq:
        seta(slide, f"s6_seta_{x:.2f}", x + w1 / 2, y1 + h1, y2)
    retangulo(slide, "s6_pi", x2, y2, w2, h2)
    t = caixa_texto(slide, "s6_pi_t", x2 + 0.15, y2 + 0.08, w2 - 0.30, 0.20, "t")
    run(paragrafo(t), "Vetor de retornos de equilíbrio implícitos", 9.5, BODY, F_REG)
    f = caixa_texto(slide, "s6_pi_f", x2 + 0.15, y2 + 0.28, w2 - 0.30, 0.32)
    p = paragrafo(f)
    run(p, "π = δ Σ w", SYM)
    run(p, "mkt", 11, sub=True)
    d = caixa_texto(slide, "s6_pi_d", x2 + w2 - 1.2, y2 + h2 - 0.24, 1.05, 0.18, "b")
    run(paragrafo(d, PP_ALIGN.RIGHT), "9 × 1", 8, MUTED, F_REG)

    # --- linha 3: distribuições
    y3, h3, w3 = 3.36, 1.62, 3.60
    cx_esq = x2 + w2 / 2
    cx_dir = (xs_dir[0] + xs_dir[1] + 1.95) / 2
    seta(slide, "s6_seta_pi", cx_esq, y2 + h2, y3)
    for x in xs_dir:
        seta(slide, f"s6_seta_{x:.2f}", x + 1.95 / 2, y1 + h1, y3)
    caixa_distribuicao(slide, "s6_prior", cx_esq - w3 / 2, y3, w3, h3,
                       "Distribuição a priori de equilíbrio",
                       [("N ~ ", SYM, TEXT, False), ("( π , τΣ )", SYM, TEXT, False)],
                       BODY, rodape="τ = 1/504")
    caixa_distribuicao(slide, "s6_view", cx_dir - w3 / 2, y3, w3, h3,
                       "Distribuição das views",
                       [("N ~ ( ", SYM, TEXT, False), ("Q", SYM, GOLD, False), (" , ", SYM, TEXT, False),
                        ("Ω", SYM, GOLD, False), (" )", SYM, TEXT, False)],
                       GOLD, rodape="k views · Q e Ω vêm do Polymarket", sigma=0.13)

    # --- linha 4: distribuição combinada
    y4, h4 = 5.30, 1.52
    x4, w4 = 3.00, 8.00
    seta(slide, "s6_seta_prior", cx_esq, y3 + h3, y4)
    seta(slide, "s6_seta_view", cx_dir, y3 + h3, y4)
    retangulo(slide, "s6_post", x4, y4, w4, h4)
    t = caixa_texto(slide, "s6_post_t", x4 + 0.15, y4 + 0.10, w4 - 0.30, 0.22, "t")
    run(paragrafo(t), "Nova distribuição combinada de retornos", 10, BODY, F_REG)
    sininha(slide, "s6_post_curva", x4 + w4 / 2 - 1.6, y4 + 0.38, 3.2, 0.42, TEXT, 2.0, sigma=0.15)
    f = caixa_texto(slide, "s6_post_f", x4 + 0.15, y4 + 0.90, w4 - 0.30, 0.32)
    p = paragrafo(f)
    run(p, "N ~ ( E[R] ,  [ (τΣ)⁻¹ + ", SYM)
    run(p, "P", SYM, GOLD); run(p, "ᵀ", SYM); run(p, "Ω", SYM, GOLD); run(p, "⁻¹", SYM); run(p, "P", SYM, GOLD)
    run(p, " ]⁻¹ )", SYM)
    e = caixa_texto(slide, "s6_post_e", x4 + 0.15, y4 + h4 - 0.34, w4 - 0.30, 0.26, "b")
    p = paragrafo(e)
    run(p, "E[R] = [ (τΣ)⁻¹ + PᵀΩ⁻¹P ]⁻¹ [ (τΣ)⁻¹ π + PᵀΩ⁻¹ ", 10.5, BODY, F_REG)
    run(p, "Q", 10.5, GOLD, F_REG); run(p, " ]", 10.5, BODY, F_REG)
    run(p, "     →   w = (δΣ)⁻¹ E[R]", 10.5, MUTED, F_REG)

    # fonte
    src = caixa_texto(slide, "s6_fonte", L, 6.88, R - L, 0.18, "b")
    run(paragrafo(src, PP_ALIGN.RIGHT), "Figura adaptada de Idzorek (2005), A step-by-step guide to the Black-Litterman model.", 7.5, FOOT, F_REG)

    if slide.has_notes_slide:
        slide.notes_slide.notes_text_frame.text = (
            "Diagrama de Idzorek na nossa notação. Esquerda: o que o mercado já carrega (δ, Σ, w_mkt → π, "
            "distribuição a priori). Direita: o que o Polymarket diz (P, Q, Ω → distribuição das views). "
            "Embaixo: o BL junta as duas — a média combinada E[R] vira a carteira do dia. "
            "O slide seguinte abre como Q e Ω são calculados.")


def render(arquivo, png):
    ps = (f"$app = New-Object -ComObject PowerPoint.Application; "
          f"$p = $app.Presentations.Open('{arquivo}', $true, $false, $false); "
          f"$p.Slides.Item({SLIDE}).Export('{png}', 'PNG', 1920, 1080); $p.Close(); $app.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)


if __name__ == "__main__":
    prs = Presentation(ARQUIVO)
    refazer(prs.slides[SLIDE - 1])
    prs.save(ARQUIVO)
    print("ok", ARQUIVO.relative_to(RAIZ), f"slide {SLIDE}")
    if "--render" in sys.argv:
        png = Path(os.environ.get("RENDER_DIR", RAIZ / "Final")) / f"slide{SLIDE:02d}_t6.png"
        render(ARQUIVO, png)
        print("PNG em", png)

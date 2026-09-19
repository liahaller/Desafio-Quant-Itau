"""Templates de reestilização do deck da final — `Final/Reestilizar/`.

Seis propostas de estilo, cada uma com UM slide de exemplo com o MESMO
conteúdo (título, caixas com seta, texto, fórmula nativa do PowerPoint,
gráfico nativo, caixa de conclusão, rodapé/navegação), para a comparação
ser só de estilo. Gradiente: do mais parecido com o deck atual até o mais
parecido com os decks de challenge (Polaris/Insper, Wolves/Insper,
Waterloo/Rotman 2026) que estão na mesma pasta.

  T1  Kairós sóbrio      escuro   mesmo deck, sem enfeites (chapado, sem glow, sem número fantasma)
  T2  Kairós claro       claro    mesma estrutura em fundo branco, dourado escurecido
  T3  Terminal           escuro   formato research (Seção | Título, nav inferior, círculos numerados) em grafite
  T4  Consultoria        claro    Polaris: cinza-claro, painéis brancos com barra dourada, aba lateral
  T5  Research           claro    Waterloo: branco, barras pretas, acento vermelho, caixa tracejada
  T6  Terminal Polymarket escuro  T3 com o azul do Polymarket só no gráfico

Só cria os templates; não toca em nenhum outro .pptx. Dados do exemplo são
ilustrativos (não são números da entrega).

Uso:
    python scripts/templates_reestilizar.py            # gera só o T6 (escolhido)
    python scripts/templates_reestilizar.py --todos    # regera os seis
    python scripts/templates_reestilizar.py --render   # e exporta PNG via PowerPoint (COM)
"""

import copy
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
PASTA = RAIZ / "Final" / "Reestilizar"

MC = 'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
A14 = 'xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main"'
M = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'

NAV = ["Motivação", "Dados", "Modelo", "Backtest", "Análise", "Apêndice"]
ATIVA = "Modelo"

# ------------------------------------------------------------------ estilos
# Chaves comuns; o layout lê `header`, `nav`, `vtab`, `badge`, `takeaway`, `arrow`.
ESTILOS = {
    "T1_kairos_sobrio_escuro": dict(
        nome="T1 · Kairós sóbrio (escuro)",
        bg="0B1426", text="F5F7FB", body="BAC6DA", muted="8395B5", rule="2A3F63",
        accent="FFB531", accent_txt="0B1426", accent2="4F8EF7",
        panel="121E36", panel_line="2A3F63", grid="1E2E4A", bar1="4F8EF7", bar2="FFB531",
        f_title="Bahnschrift SemiBold Condensed", f_body="Bahnschrift SemiLight",
        f_kick="Bahnschrift SemiBold", f_num="Bahnschrift SemiBold Condensed",
        header="kairos", nav="footer", vtab=False, badge=False, takeaway="bar", arrow="line",
    ),
    "T2_kairos_claro": dict(
        nome="T2 · Kairós claro",
        bg="FFFFFF", text="111827", body="374151", muted="6B7280", rule="D1D5DB",
        accent="C98A12", accent_txt="FFFFFF", accent2="2F6FD6",
        panel="F7F8FA", panel_line="D1D5DB", grid="E5E7EB", bar1="2F6FD6", bar2="C98A12",
        f_title="Bahnschrift SemiBold Condensed", f_body="Bahnschrift Light",
        f_kick="Bahnschrift SemiBold", f_num="Bahnschrift SemiBold Condensed",
        header="kairos", nav="footer", vtab=False, badge=False, takeaway="bar", arrow="line",
    ),
    "T3_terminal_escuro": dict(
        nome="T3 · Terminal (escuro, formato research)",
        bg="15181D", text="FFFFFF", body="D5D9E0", muted="9AA3AF", rule="3A404A",
        accent="FFB531", accent_txt="15181D", accent2="9AA3AF",
        panel="1E2229", panel_line="3A404A", grid="2A2F37", bar1="6B7280", bar2="FFB531",
        f_title="Segoe UI Semibold", f_body="Segoe UI", f_kick="Segoe UI Semibold",
        f_num="Segoe UI Semibold",
        header="pipe", nav="bar", vtab=False, badge=True, takeaway="dashed", arrow="line",
    ),
    "T4_consultoria_claro": dict(
        nome="T4 · Consultoria (claro, Polaris)",
        bg="F2F2F2", text="1F1F1F", body="3F3F3F", muted="7A7A7A", rule="C8C8C8",
        accent="FFB531", accent_txt="1F1F1F", accent2="3A3A3A",
        panel="FFFFFF", panel_line="DDDDDD", grid="E3E3E3", bar1="3A3A3A", bar2="FFB531",
        f_title="Segoe UI Semibold", f_body="Segoe UI", f_kick="Segoe UI Semibold",
        f_num="Segoe UI Semibold",
        header="sub", nav="bar", vtab=True, badge=True, takeaway="dark", arrow="line",
    ),
    "T5_research_claro": dict(
        nome="T5 · Research (claro, Waterloo)",
        bg="FFFFFF", text="1A1A1A", body="333333", muted="666666", rule="1A1A1A",
        accent="D62828", accent_txt="FFFFFF", accent2="2B2B2B",
        panel="F2F2F2", panel_line="F2F2F2", grid="E0E0E0", bar1="2B2B2B", bar2="D62828",
        f_title="Segoe UI Semibold", f_body="Segoe UI", f_kick="Segoe UI Semibold",
        f_num="Segoe UI Semibold",
        header="pipe", nav="bar", vtab=False, badge=True, takeaway="dashed", arrow="block",
    ),
    "T6_terminal_polymarket": dict(          # T3 com a série de mercado do gráfico em azul Polymarket
        nome="T6 · Terminal Polymarket (escuro, azul)",
        bg="15181D", text="FFFFFF", body="D5D9E0", muted="9AA3AF", rule="3A404A",
        accent="FFB531", accent_txt="15181D", accent2="2A4BC4",
        panel="1E2229", panel_line="3A404A", grid="2A2F37", bar1="2A4BC4", bar2="FFB531",
        f_title="Segoe UI Semibold", f_body="Segoe UI", f_kick="Segoe UI Semibold",
        f_num="Segoe UI Semibold",
        header="pipe", nav="bar", vtab=False, badge=True, takeaway="dashed", arrow="line",
    ),
}


# --------------------------------------------------------------- primitivas
def rgb(h):
    return RGBColor.from_string(h)


def tx(slide, x, y, w, h, runs, font, pt, color, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, spc=None, italic=False, line_sp=None):
    """Caixa de texto; `runs` = str ou lista de (texto, negrito, subscrito)."""
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if line_sp:
        p.line_spacing = line_sp
    if isinstance(runs, str):
        runs = [(runs, False, False)]
    for texto, negrito, sub in runs:
        r = p.add_run()
        r.text = texto
        f = r.font
        f.name, f.size, f.bold, f.italic = font, Pt(pt), negrito, italic
        f.color.rgb = rgb(color)
        if spc:
            r._r.get_or_add_rPr().set("spc", str(spc))
        if sub:
            r._r.get_or_add_rPr().set("baseline", "-25000")
    return sh


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, dash=False,
         shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = MSO_LINE.DASH
    else:
        sh.line.fill.background()
    sh.text_frame.text = ""
    return sh


def arrow_line(slide, x1, y1, x2, y2, color, lw=1.25):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    ln = c.line._get_or_add_ln()
    ln.append(parse_xml(f'<a:tailEnd {nsdecls("a")} type="triangle" w="med" len="med"/>'))
    return c


def badge(slide, cx, cy, d, n, st):
    sh = rect(slide, cx - d / 2, cy - d / 2, d, d, fill=st["accent"], shape=MSO_SHAPE.OVAL)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.name, r.font.size, r.font.bold = st["f_kick"], Pt(10), True
    r.font.color.rgb = rgb(st["accent_txt"])
    return sh


# --------------------------------------------------------------- fórmula
def _mr(t, st, sz):
    return (f'<m:r><a:rPr lang="en-US" sz="{sz}"><a:solidFill>'
            f'<a:srgbClr val="{st["text"]}"/></a:solidFill></a:rPr><m:t>{t}</m:t></m:r>')


def _sup(e, s):
    return f"<m:sSup><m:e>{e}</m:e><m:sup>{s}</m:sup></m:sSup>"


def _sub(e, s):
    return f"<m:sSub><m:e>{e}</m:e><m:sub>{s}</m:sub></m:sSub>"


def _d(e, beg="(", end=")"):
    return (f'<m:d><m:dPr><m:begChr m:val="{beg}"/><m:endChr m:val="{end}"/></m:dPr>'
            f"<m:e>{e}</m:e></m:d>")


def formula(slide, x, y, w, h, st, pt=20):
    """μ_BL = [(τΣ)⁻¹ + PᵀΩ⁻¹P]⁻¹ [(τΣ)⁻¹π + PᵀΩ⁻¹Q] como equação NATIVA (OMML).

    O PowerPoint 2010+ lê o ramo `a14:m`; leitores antigos caem no texto do
    `mc:Fallback`. A equação fica editável no PowerPoint (Inserir > Equação).
    """
    sz = pt * 100
    r = lambda t: _mr(t, st, sz)  # noqa: E731
    ts_inv = _sup(_d(r("τΣ")), r("−1"))
    pt_om = _sup(r("P"), r("T")) + _sup(r("Ω"), r("−1"))
    inner1 = ts_inv + r("+") + pt_om + r("P")
    inner2 = ts_inv + r("π") + r("+") + pt_om + r("Q")
    math = (_sub(r("μ"), r("BL")) + r("=") + _sup(_d(inner1, "[", "]"), r("−1"))
            + _d(inner2, "[", "]"))
    p_math = parse_xml(
        f'<a:p {nsdecls("a")} {A14}><a14:m><m:oMathPara {M}><m:oMathParaPr>'
        f'<m:jc m:val="left"/></m:oMathParaPr><m:oMath>{math}</m:oMath>'
        f"</m:oMathPara></a14:m></a:p>")

    fallback = tx(slide, x, y, w, h, "μ_BL = [(τΣ)^-1 + P'Ω^-1 P]^-1 [(τΣ)^-1 π + P'Ω^-1 Q]",
                  "Cambria Math", pt, st["text"], anchor=MSO_ANCHOR.MIDDLE)
    sp_fb = fallback._element
    sp = copy.deepcopy(sp_fb)
    tx_body = sp.find(qn("p:txBody"))
    tx_body.replace(tx_body.find(qn("a:p")), p_math)

    ac = parse_xml(f'<mc:AlternateContent {MC}><mc:Choice Requires="a14" {A14}/>'
                   f"<mc:Fallback/></mc:AlternateContent>")
    parent = sp_fb.getparent()
    i = parent.index(sp_fb)
    parent.remove(sp_fb)
    ac[0].append(sp)
    ac[1].append(sp_fb)
    parent.insert(i, ac)


# --------------------------------------------------------------- gráfico
def grafico(slide, x, y, w, h, st):
    """Colunas agrupadas NATIVAS (editáveis no PowerPoint). Dados ilustrativos."""
    cd = CategoryChartData()
    cd.categories = ["XLU", "TLT", "TIP", "XLK", "XLE", "SPY"]
    cd.add_series("peso de mercado", (0.12, 0.20, 0.10, 0.28, 0.08, 0.22))
    cd.add_series("peso Black-Litterman", (0.16, 0.24, 0.13, 0.20, 0.08, 0.19))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y),
                                Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_title = False
    ch.font.name, ch.font.size = st["f_body"], Pt(9)
    ch.font.color.rgb = rgb(st["muted"])
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    plot = ch.plots[0]
    plot.gap_width, plot.overlap = 70, -10
    for ser, cor in zip(plot.series, (st["bar1"], st["bar2"])):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(cor)
        ser.format.line.fill.background()
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(st["grid"])
    va.major_gridlines.format.line.width = Pt(0.5)
    va.format.line.fill.background()
    va.tick_labels.number_format = "0%"
    va.tick_labels.number_format_is_linked = False
    va.maximum_scale, va.major_unit = 0.30, 0.10
    ca = ch.category_axis
    ca.format.line.color.rgb = rgb(st["rule"])
    ca.format.line.width = Pt(0.75)
    # fundo transparente do gráfico e da área de plotagem
    cs = ch._chartSpace
    if cs.find(qn("c:roundedCorners")) is None:
        cs.insert(0, parse_xml(f'<c:roundedCorners {nsdecls("c")} val="0"/>'))
    c_chart = cs.find(qn("c:chart"))
    c_chart.addnext(parse_xml(
        f'<c:spPr {nsdecls("c")} {nsdecls("a")}><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'))
    cs.chart.plotArea.append(parse_xml(f'<c:spPr {nsdecls("c")} {nsdecls("a")}><a:noFill/></c:spPr>'))
    return gf


# --------------------------------------------------------------- blocos
def cabecalho(slide, st):
    W = 13.333
    if st["header"] == "kairos":            # deck atual: KAIROS | seção, título grande
        rect(slide, 0.80, 0.50, 0.06, 0.22, fill=st["accent"])
        tx(slide, 0.95, 0.47, 1.2, 0.3, "KAIROS", st["f_kick"], 12, st["accent"], spc=300)
        tx(slide, 2.05, 0.47, 6, 0.3, "Modelo", st["f_body"], 12, st["muted"])
        tx(slide, 11.5, 0.47, 1.03, 0.3, "09", st["f_kick"], 10, st["muted"], align=PP_ALIGN.RIGHT)
        rect(slide, 0.80, 0.80, W - 1.6, 0.01, fill=st["rule"])
        tx(slide, 0.80, 0.95, 11.5, 0.6,
           "A probabilidade do Polymarket vira uma view com incerteza no Black-Litterman",
           st["f_title"], 26, st["text"])
    elif st["header"] == "pipe":            # research: Seção | Título, régua, marca à direita
        tx(slide, 0.55, 0.38, 11, 0.5,
           [("Modelo", True, False), ("  |  ", False, False),
            ("A probabilidade do Polymarket vira uma view com incerteza no BL", False, False)],
           st["f_title"], 20, st["text"], anchor=MSO_ANCHOR.MIDDLE)
        tx(slide, 11.3, 0.38, 1.5, 0.5, "KAIRÓS", st["f_kick"], 11, st["accent"],
           align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, spc=300)
        rect(slide, 0.55, 0.95, W - 1.1, 0.015, fill=st["rule"])
    else:                                   # consultoria: título + subtítulo cinza, marca à direita
        tx(slide, 0.85, 0.30, 10, 0.55,
           "A view do Polymarket no Black-Litterman", st["f_title"], 26, st["text"])
        tx(slide, 0.85, 0.85, 10, 0.3,
           "Do preço do contrato ao peso da carteira, em três passos",
           st["f_body"], 12, st["muted"])
        tx(slide, 11.3, 0.35, 1.5, 0.4, "KAIRÓS", st["f_kick"], 12, st["text"],
           align=PP_ALIGN.RIGHT, spc=300)
        rect(slide, 11.55, 0.72, 1.25, 0.04, fill=st["accent"])


def fluxo(slide, x0, y0, st):
    """Três caixas ligadas por setas — o caminho probabilidade → surpresa → Q."""
    caixas = [("POLYMARKET", "72 %", "P(corte): preço do contrato “Sim”"),
              ("SURPRESA", "+0,18", "vs. o que a curva de juros já precifica"),
              ("VIEW  Q", "+42 bps", "β × surpresa, no par XLU − XLK")]
    w, h, gap = 1.72, 1.05, 0.34
    for i, (k, num, sub) in enumerate(caixas):
        x = x0 + i * (w + gap)
        fill = st["panel"]
        line = st["panel_line"] if st["header"] != "pipe" or st["bg"] != "FFFFFF" else None
        sh = rect(slide, x, y0, w, h, fill=fill, line=line)
        if st["header"] == "sub":           # painel Polaris: barra dourada no topo
            rect(slide, x, y0, w, 0.05, fill=st["accent"])
        if st["takeaway"] == "dashed" and st["bg"] == "FFFFFF":   # Waterloo: rótulo em barra preta
            rect(slide, x, y0, w, 0.26, fill=st["accent2"])
            tx(slide, x, y0, w, 0.26, k, st["f_kick"], 8.5, "FFFFFF", align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE, spc=150)
        else:
            tx(slide, x + 0.12, y0 + 0.09, w - 0.24, 0.2, k, st["f_kick"], 8.5,
               st["accent"] if st["bg"] != "F2F2F2" else st["muted"], spc=200)
        tx(slide, x + 0.12, y0 + 0.33, w - 0.24, 0.4, num, st["f_num"],
           20 if "Bahnschrift" in st["f_num"] else 16, st["text"])
        tx(slide, x + 0.12, y0 + 0.72, w - 0.24, 0.3, sub, st["f_body"], 8, st["muted"])
        if st["badge"]:
            badge(slide, x + w - 0.02, y0 + 0.02, 0.28, i + 1, st)
        if i < 2:
            ya = y0 + h / 2
            if st["arrow"] == "block":
                a = rect(slide, x + w + 0.05, ya - 0.11, gap - 0.10, 0.22, fill=st["muted"],
                         shape=MSO_SHAPE.RIGHT_ARROW)
                a.adjustments[0] = 0.5
            else:
                arrow_line(slide, x + w + 0.04, ya, x + w + gap - 0.04, ya,
                           st["accent"] if st["bg"] not in ("F2F2F2", "FFFFFF") else st["accent2"])
        sh.name = f"caixa_{i + 1}"


def conclusao(slide, x, y, w, h, st):
    texto = [("Onde o Polymarket concorda com o mercado a surpresa é zero e a carteira fica em w",
              False, False), ("mkt", False, True),
             ("; a view só desloca o peso quando há discordância — e Ω decide quanto.", False, False)]
    if st["takeaway"] == "bar":             # painel com barra de acento à esquerda
        rect(slide, x, y, w, h, fill=st["panel"], line=st["panel_line"])
        rect(slide, x, y, 0.06, h, fill=st["accent"])
        tx(slide, x + 0.25, y, w - 0.5, h, [("Leitura.  ", True, False)] + texto,
           st["f_body"], 12, st["text"], anchor=MSO_ANCHOR.MIDDLE)
    elif st["takeaway"] == "dashed":        # Waterloo: caixa tracejada
        rect(slide, x, y, w, h, line=st["rule"], lw=0.75, dash=True)
        tx(slide, x + 0.25, y, w - 0.5, h, texto, st["f_body"], 12, st["text"],
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    else:                                   # Polaris: painel grafite com texto branco
        rect(slide, x, y, w, h, fill=st["accent2"])
        rect(slide, x, y, 0.08, h, fill=st["accent"])
        tx(slide, x + 0.3, y, w - 0.6, h, texto, st["f_body"], 12, "FFFFFF",
           anchor=MSO_ANCHOR.MIDDLE)


def rodape(slide, st):
    W, H = 13.333, 7.5
    if st["nav"] == "footer":               # deck atual
        rect(slide, 0.80, 6.88, W - 1.6, 0.01, fill=st["rule"])
        tx(slide, 0.80, 6.96, 6, 0.25, "ITAÚ QUANT AI CHALLENGE  ·  FINAL 2026",
           st["f_kick"], 8, st["muted"], spc=200)
        tx(slide, 6.5, 6.96, 6.03, 0.25, "BLACK-LITTERMAN  ×  POLYMARKET", st["f_kick"], 8,
           st["muted"], align=PP_ALIGN.RIGHT, spc=200)
        return
    # barra de navegação dos challenges: seções, ativa em destaque, nº da página
    y = H - 0.42
    rect(slide, 0.55, y - 0.08, W - 1.1, 0.01, fill=st["rule"] if st["bg"] != "FFFFFF" else "BBBBBB")
    n = len(NAV)
    x0, wtot = 0.55, W - 1.1 - 0.8
    for i, item in enumerate(NAV):
        xi = x0 + i * wtot / n
        ativa = item == ATIVA
        tx(slide, xi, y, wtot / n, 0.3, item.upper() if st["vtab"] else item,
           st["f_kick"] if ativa else st["f_body"], 8.5,
           st["text"] if ativa else st["muted"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if ativa:
            rect(slide, xi + wtot / n / 2 - 0.45, y + 0.29, 0.9, 0.03, fill=st["accent"])
    tx(slide, W - 1.2, y, 0.65, 0.3, "9", st["f_body"], 9, st["muted"], align=PP_ALIGN.RIGHT,
       anchor=MSO_ANCHOR.MIDDLE)


def aba_vertical(slide, st):
    """Polaris: barra grafite à esquerda com o nome da seção rotacionado."""
    rect(slide, 0, 0, 0.36, 7.5, fill=st["accent2"])
    sh = tx(slide, 0.18 - 2.0, 3.6 - 0.18, 4.0, 0.36, "Modelo  ·  Black-Litterman", st["f_body"],
            12, "FFFFFF", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spc=150)
    sh.rotation = 270


# --------------------------------------------------------------- slide
def monta(chave, st):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(st["bg"])

    if st["vtab"]:
        aba_vertical(s, st)
    cabecalho(s, st)
    L = 0.85 if st["vtab"] else (0.80 if st["header"] == "kairos" else 0.55)
    R = 12.53 if st["header"] == "kairos" else 12.78
    yt = {"kairos": 1.75, "pipe": 1.40, "sub": 1.45}[st["header"]]   # topo do conteúdo

    # coluna esquerda: kicker, fluxo, texto, fórmula
    kick_cor = st["accent"] if st["bg"] not in ("F2F2F2",) else st["muted"]
    tx(s, L, yt, 6, 0.25, "COMO A VIEW NASCE", st["f_kick"], 9, kick_cor, spc=200)
    fluxo(s, L, yt + 0.35, st)
    tx(s, L, yt + 1.6, 5.9, 0.75,
       [("A view entra na mesma unidade do prior de equilíbrio ", False, False),
        ("π", True, False), (": retorno esperado do par. ", False, False),
        ("A incerteza ", False, False), ("Ω", True, False),
        (" decide quanto o peso se afasta de w", False, False), ("mkt", False, True),
        ("; com Ω grande, o posterior volta ao prior.", False, False)],
       st["f_body"], 11.5, st["body"], line_sp=1.15)
    tx(s, L, yt + 2.5, 6, 0.25, "POSTERIOR DE BLACK-LITTERMAN", st["f_kick"], 9, kick_cor,
       spc=200)
    formula(s, L, yt + 2.8, 5.9, 0.65, st, pt=20)
    tx(s, L, yt + 3.5, 5.9, 0.25,
       "Q: view do Polymarket  ·  Ω: incerteza da view  ·  τ: escala do prior",
       st["f_body"], 9, st["muted"])

    # coluna direita: kicker + gráfico nativo
    xg = 7.15
    tx(s, xg, yt, 5.5, 0.25, "PESO DE MERCADO × PESO BL  (EXEMPLO)", st["f_kick"], 9, kick_cor,
       spc=200)
    if st["header"] == "sub":               # painel branco atrás do gráfico (Polaris)
        rect(s, xg - 0.15, yt + 0.3, R - xg + 0.15, 3.55, fill=st["panel"], line=st["panel_line"])
        rect(s, xg - 0.15, yt + 0.3, R - xg + 0.15, 0.05, fill=st["accent"])
    grafico(s, xg, yt + 0.4, R - xg, 3.4, st)

    # conclusão + fonte
    conclusao(s, L, yt + 4.05, R - L, 0.62, st)
    tx(s, L, yt + 4.78, R - L, 0.25,
       "Fonte: elaboração própria. Dados ilustrativos — template de estilo, não são números da entrega.",
       st["f_body"], 8, st["muted"], italic=True)
    rodape(s, st)

    s.notes_slide.notes_text_frame.text = st["nome"]
    out = PASTA / f"{chave}.pptx"
    prs.save(out)
    return out


def render(arquivos):
    """Exporta cada template em PNG pelo PowerPoint (COM) — só para conferir."""
    linhas = ["$app = New-Object -ComObject PowerPoint.Application"]
    for f in arquivos:
        png = str(f.with_suffix(".png")).replace("/", "\\")
        linhas += [f"$p = $app.Presentations.Open('{f}', $true, $false, $false)",
                   f"$p.Slides.Item(1).Export('{png}', 'PNG', 1920, 1080)", "$p.Close()"]
    linhas.append("$app.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", "; ".join(linhas)], check=True)


if __name__ == "__main__":
    PASTA.mkdir(exist_ok=True)
    # T1–T5 foram descartados pelo dono (2026-09-18); ficam no dicionário e só
    # saem com --todos. Por padrão gera apenas o T6.
    chaves = ESTILOS if "--todos" in sys.argv else ["T6_terminal_polymarket"]
    saidas = [monta(k, ESTILOS[k]) for k in chaves]
    for f in saidas:
        print("ok", f.relative_to(RAIZ))
    if "--render" in sys.argv:
        render(saidas)
        print("PNG exportados")

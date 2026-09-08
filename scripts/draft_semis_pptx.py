"""Wireframe do deck da semifinal (draft, nao e versao final).

Gera `Semis/Draft_Semis_v1.pptx` a partir do `Semis/draft_apresentacao.txt`:
foco em COMPOSICAO (onde cada coisa fica, que itens aparecem escritos), nao em
conteudo final. Nenhum grafico e gerado — todo visual que dependeria de figura
vira uma caixa tracejada dizendo o que iria ali.

Regra de texto: itens curtos que guiam a fala, nunca paragrafo. O `--demo`
falha se algum slide passar do teto de palavras ou alguma linha ficar longa
demais para um slide falado por cima.

Paleta herdada de `scripts/relatorio_p4.py` (mesma identidade do relatorio).
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
SAIDA = RAIZ / "Semis" / "Draft_Semis_v1.pptx"

# paleta (igual a do relatorio, para o draft ja nascer na identidade certa)
BG = RGBColor(0x0A, 0x16, 0x21)
RULE = RGBColor(0x24, 0x37, 0x4A)
ORANGE = RGBColor(0xF5, 0xA6, 0x23)
TEXT = RGBColor(0xE8, 0xED, 0xF2)
MUTED = RGBColor(0x6B, 0x82, 0x99)
DIM = RGBColor(0x39, 0x4E, 0x63)
FONTE = "Segoe UI"

L, R = 0.62, 12.71          # margens de coluna, em polegadas
TOPO, BASE = 0.95, 6.92     # reguas horizontais

# tetos de texto por slide — sao a regra "sem paragrafo", verificada no --demo
MAX_PALAVRAS = 180       # rede de seguranca por slide (rotulo curto e barato)
MAX_PALAVRAS_LINHA = 12  # acima disso a linha virou frase corrida
MAX_FRASES = 6           # o criterio de verdade: linhas com >= 8 palavras
MAX_LINHA = 64


# --------------------------------------------------------------------- utils
def txt(slide, x, y, w, h, texto, tam=11, cor=TEXT, negrito=False,
        alinha=PP_ALIGN.LEFT, espaco=1.15, maiuscula=False, ancora=MSO_ANCHOR.TOP):
    """Caixa de texto sem moldura. `texto` pode ser str ou lista de linhas."""
    cx = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = cx.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancora
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    linhas = [texto] if isinstance(texto, str) else list(texto)
    for i, linha in enumerate(linhas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinha
        p.line_spacing = espaco
        run = p.add_run()
        run.text = linha.upper() if maiuscula else linha
        run.font.size = Pt(tam)
        run.font.bold = negrito
        run.font.color.rgb = cor
        run.font.name = FONTE
    return cx


def caixa(slide, x, y, w, h, cor_linha=RULE, preenche=None, tracejada=False,
          espessura=1.0):
    """Retangulo de moldura (o tijolo do wireframe)."""
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    if preenche is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = preenche
    sh.line.color.rgb = cor_linha
    sh.line.width = Pt(espessura)
    if tracejada:
        sh.line.dash_style = 4  # MSO_LINE_DASH_STYLE.DASH
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def placeholder(slide, x, y, w, h, tipo, descricao, itens=()):
    """Caixa tracejada no lugar de um grafico/imagem que NAO foi gerado."""
    caixa(slide, x, y, w, h, cor_linha=DIM, tracejada=True)
    txt(slide, x + 0.22, y + 0.20, w - 0.44, 0.28, f"[ {tipo} ]", tam=9.5,
        cor=ORANGE, negrito=True, maiuscula=True)
    txt(slide, x + 0.22, y + 0.52, w - 0.44, 0.6, descricao, tam=11, cor=TEXT)
    if itens:
        txt(slide, x + 0.22, y + h - 0.30 - 0.20 * len(itens), w - 0.44,
            0.22 * len(itens), [f"· {i}" for i in itens], tam=9, cor=MUTED)
    return None


def cabecalho(slide, numero, titulo, tempo, sobretitulo=None):
    txt(slide, L, 0.34, 6.0, 0.3, "KAIROS", tam=10, cor=ORANGE, negrito=True)
    rot = f"{sobretitulo} · " if sobretitulo else ""
    txt(slide, L + 1.25, 0.34, 8.4, 0.3, f"{rot}{titulo}", tam=10, cor=MUTED,
        maiuscula=True)
    txt(slide, R - 2.6, 0.34, 2.6, 0.3, f"{numero}  ·  {tempo}", tam=10,
        cor=MUTED, alinha=PP_ALIGN.RIGHT)
    regua(slide, L, TOPO, R)


def regua(slide, x1, y, x2, cor=RULE):
    ln = slide.shapes.add_connector(1, Inches(x1), Inches(y), Inches(x2), Inches(y))
    ln.line.color.rgb = cor
    ln.line.width = Pt(0.75)
    return ln


def rodape(slide, texto):
    """Nota de producao: vai para as ANOTACOES do slide, nao para a tela.

    O draft e falado por cima; nota que aparece projetada e texto a mais.
    """
    slide.notes_slide.notes_text_frame.text = texto


def seta(slide, x1, y1, x2, y2, cor=DIM):
    ln = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = cor
    ln.line.width = Pt(1.25)
    return ln


def chips(slide, x, y, w, h, itens, gap=0.16, cor=RULE, cor_txt=TEXT, tam=10):
    """Fileira horizontal de caixinhas curtas."""
    n = len(itens)
    cw = (w - gap * (n - 1)) / n
    for i, item in enumerate(itens):
        cx = x + i * (cw + gap)
        caixa(slide, cx, y, cw, h, cor_linha=cor)
        txt(slide, cx + 0.12, y + (h - 0.24) / 2, cw - 0.24, 0.3, item, tam=tam,
            cor=cor_txt, alinha=PP_ALIGN.CENTER)
    return cw


def linhas_tabela(slide, x, y, w, cabec, linhas, tam=9.5, alt=0.30, cols=None):
    """Tabela leve, so com regua sob o cabecalho."""
    cols = cols or [1 / len(cabec)] * len(cabec)
    xs, acc = [], 0.0
    for frac in cols:
        xs.append(x + acc * w)
        acc += frac
    for i, c in enumerate(cabec):
        txt(slide, xs[i], y, w * cols[i] - 0.1, 0.24, c, tam=8, cor=MUTED,
            maiuscula=True)
    regua(slide, x, y + 0.26, x + w)
    for j, linha in enumerate(linhas):
        yy = y + 0.36 + j * alt
        for i, cel in enumerate(linha):
            cor = TEXT if i == 0 else MUTED
            txt(slide, xs[i], yy, w * cols[i] - 0.1, 0.26, cel, tam=tam, cor=cor)
    return y + 0.36 + len(linhas) * alt


def slide_novo(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo = s.background.fill
    fundo.solid()
    fundo.fore_color.rgb = BG
    return s


# ------------------------------------------------------- diagrama da arquitetura
# Caixas do diagrama, em polegadas: (x, y, w, h)
D_EQ = (L, 1.35, 2.55, 1.05)
D_Q = (3.52, 1.35, 2.55, 1.05)
D_OM = (6.44, 1.35, 2.55, 1.05)
D_TAT = (L, 3.05, 8.37, 0.95)
D_TETO = (9.75, 1.35, 2.96, 1.05)
D_CART = (9.75, 3.05, 2.96, 0.95)


def diagrama(slide, ativo):
    """Desenha o diagrama da arquitetura realcando a peca do build atual.

    `ativo` ∈ {"bl", "q", "omega", "tatica"}. O que nao esta ativo fica DIM —
    nada some entre um build e o outro.
    """
    ativas = {
        "bl": {"eq", "q", "om"},
        "q": {"q"},
        "omega": {"om"},
        "tatica": {"tat", "teto"},
    }[ativo]

    def peca(nome, geo, titulo, sub, forte):
        x, y, w, h = geo
        linha = ORANGE if forte else RULE
        caixa(slide, x, y, w, h, cor_linha=linha,
              espessura=1.5 if forte else 0.75)
        txt(slide, x + 0.16, y + 0.16, w - 0.32, 0.3, titulo, tam=12,
            cor=TEXT if forte else DIM, negrito=True)
        txt(slide, x + 0.16, y + 0.52, w - 0.32, 0.4, sub, tam=8.5,
            cor=MUTED if forte else DIM)

    peca("eq", D_EQ, "EQUILÍBRIO", "o que o mercado já espera", "eq" in ativas)
    peca("q", D_Q, "OPINIÃO  (Q)", "quanto se espera render", "q" in ativas)
    peca("om", D_OM, "CONFIANÇA  (Ω)", "o quanto ela pesa", "om" in ativas)
    peca("tat", D_TAT, "CAMADA TÁTICA", "lê o movimento, não o nível",
         "tat" in ativas)
    peca("teto", D_TETO, "TETO DE RISCO", "Σ|w − w_mkt| ≤ 1", "teto" in ativas)
    peca("cart", D_CART, "CARTEIRA DO DIA", "9 ETFs", False)

    # operadores entre as tres primeiras caixas
    txt(slide, 3.12, 1.72, 0.4, 0.3, "+", tam=15, cor=DIM, alinha=PP_ALIGN.CENTER)
    txt(slide, 6.04, 1.72, 0.4, 0.3, "×", tam=15, cor=DIM, alinha=PP_ALIGN.CENTER)

    # setas para o teto e do teto para a carteira
    seta(slide, D_OM[0] + D_OM[2], 1.88, D_TETO[0], 1.88,
         ORANGE if ativo == "tatica" else DIM)
    seta(slide, D_TAT[0] + D_TAT[2], 3.52, D_TETO[0] + 0.4, 3.52,
         ORANGE if ativo == "tatica" else DIM)
    seta(slide, D_TETO[0] + 1.48, D_TETO[1] + D_TETO[3], D_TETO[0] + 1.48, D_CART[1],
         DIM)

    # rotulo "gestor" riscado sob Q e Ω (a animacao do build 1)
    for geo in (D_Q, D_OM):
        x, y, w, h = geo
        forte = ativo in ("bl", "q", "omega")
        txt(slide, x + 0.16, y + h + 0.10, 1.1, 0.24, "gestor", tam=9,
            cor=DIM if forte else DIM)
        ln = slide.shapes.add_connector(
            1, Inches(x + 0.14), Inches(y + h + 0.19),
            Inches(x + 0.68), Inches(y + h + 0.19))
        ln.line.color.rgb = ORANGE if forte else DIM
        ln.line.width = Pt(1.0)
        txt(slide, x + 0.80, y + h + 0.10, 1.7, 0.24, "POLYMARKET", tam=9,
            cor=ORANGE if forte else DIM, negrito=True)


# ------------------------------------------------------------------- os slides
def s1_capa(prs):
    s = slide_novo(prs)
    cabecalho(s, "01", "identidade e hipótese", "0:32")
    txt(s, L, 1.45, 7.6, 1.0, "KAIROS", tam=54, cor=TEXT, negrito=True)
    txt(s, L, 2.55, 7.6, 0.4, "Black-Litterman movido a probabilidade negociada",
        tam=14, cor=ORANGE)
    regua(s, L, 3.25, 8.6)
    txt(s, L, 3.50, 7.9, 1.4, [
        "O mercado de previsão reprecifica",
        "notícia antes da bolsa.",
        "Usamos esse preço como a view —",
        "no lugar do palpite do gestor.",
    ], tam=17, cor=TEXT, espaco=1.30)
    txt(s, L, 5.30, 7.9, 0.3, "kairós  ·  o instante certo de agir", tam=11,
        cor=MUTED)

    caixa(s, 9.3, 1.45, 3.41, 3.55, cor_linha=RULE)
    txt(s, 9.52, 1.65, 3.0, 0.24, "ficha técnica", tam=8, cor=ORANGE,
        maiuscula=True)
    txt(s, 9.52, 2.00, 3.0, 2.6, [
        "9 ETFs (EUA)",
        "setores · renda fixa · índice",
        "",
        "benchmark  SPY",
        "rebalanceamento  diário",
        "janela  374 pregões",
        "fev/2025 – ago/2026",
    ], tam=11, cor=TEXT, espaco=1.45)
    txt(s, 9.52, 4.45, 3.0, 0.5, ["US$ 26 bi negociados no 1ºT",
                                  "90% de acerto a 1 mês"], tam=9, cor=MUTED)
    rodape(s, "capa · cartão de texto: exceção deliberada — a banca olha para "
              "quem fala, não para a tela")
    return s


def s_build(prs, numero, titulo, tempo, ativo, detalhe):
    s = slide_novo(prs)
    cabecalho(s, numero, titulo, tempo, sobretitulo="a arquitetura")
    diagrama(s, ativo)
    regua(s, L, 4.42, R)
    detalhe(s)
    return s


def s2_bl(prs):
    def detalhe(s):
        txt(s, L, 4.62, 5.6, 0.3, "o que o modelo faz", tam=8, cor=ORANGE,
            maiuscula=True)
        txt(s, L, 4.98, 5.9, 1.6, [
            "parte do equilíbrio, não do palpite",
            "só se afasta quando recebe uma opinião",
            "confiança zero  →  fica exatamente o mercado",
        ], tam=13, cor=TEXT, espaco=1.55)
        txt(s, 7.3, 4.62, 5.4, 0.3, "a virada deste slide", tam=8, cor=ORANGE,
            maiuscula=True)
        txt(s, 7.3, 4.98, 5.4, 1.6, [
            "Q e Ω são digitados pelo gestor",
            "→ os dois passam a vir do Polymarket",
            "nenhum dos dois é declarado por nós",
        ], tam=13, cor=TEXT, espaco=1.55)
        rodape(s, "animação: ao falar 'digitados pelo gestor', riscar o rótulo "
                  "e entrar POLYMARKET sob Q e Ω")
    return s_build(prs, "02", "1 de 4 — como o Black-Litterman funciona", "0:22",
                   "bl", detalhe)


def s3_q(prs):
    def detalhe(s):
        txt(s, L, 4.62, 6.0, 0.3, "de onde vem o Q", tam=8, cor=ORANGE,
            maiuscula=True)
        chips(s, L, 4.95, 8.37, 0.52,
              ["probabilidade", "− âncora", "surpresa (bps)", "× β medido",
               "retorno / ativo"], tam=9)
        txt(s, L, 5.60, 8.37, 0.28, "Q = (E_poly − âncora) · Σ Pᵢ βᵢ", tam=10,
            cor=MUTED)
        txt(s, L, 5.98, 8.37, 0.28, "β medido em anúncios passados — não arbitrado",
            tam=11, cor=TEXT)
        txt(s, 9.3, 4.62, 3.4, 0.3, "as 4 views", tam=8, cor=ORANGE,
            maiuscula=True)
        linhas_tabela(s, 9.3, 4.95, 3.41, ["view", "ativa", "exp."], [
            ["2.3  Fed", "312/374", "neutra"],
            ["2.2  inflação", "253/374", "neutra"],
            ["B  taxa fim do ano", "154/374", "neutra"],
            ["15b  véspera", "26/374", "direcional"],
        ], tam=8.5, alt=0.26, cols=[0.46, 0.28, 0.26])
        rodape(s, "a fala percorre UMA view inteira (o Fed) — a tabela existe "
                  "para as outras três serem só nomeadas")
    return s_build(prs, "03", "2 de 4 — de onde vem o Q", "0:41", "q", detalhe)


def s4_omega(prs):
    def detalhe(s):
        txt(s, L, 4.62, 6.0, 0.3, "de onde vem o Ω", tam=8, cor=ORANGE,
            maiuscula=True)
        for i, (t, sub) in enumerate([
            ("ESTABILIDADE", "quanto a distribuição se moveu"),
            ("COERÊNCIA", "o quanto o livro deixa de somar 1"),
        ]):
            x = L + i * 2.35
            caixa(s, x, 4.95, 2.15, 0.85, cor_linha=RULE)
            txt(s, x + 0.14, 5.08, 1.9, 0.26, t, tam=10, cor=TEXT, negrito=True)
            txt(s, x + 0.14, 5.36, 1.9, 0.36, sub, tam=8, cor=MUTED)
        caixa(s, L + 4.85, 4.95, 3.52, 0.85, cor_linha=ORANGE, espessura=1.5)
        txt(s, L + 5.00, 5.08, 3.2, 0.26, "c ≥ 1", tam=13, cor=ORANGE,
            negrito=True)
        txt(s, L + 5.00, 5.38, 3.2, 0.3, "o BL clássico é o teto, nunca o piso",
            tam=9, cor=TEXT)
        txt(s, L, 6.02, 8.37, 0.6, ["calibrada contra o erro da probabilidade",
                                    "nunca contra o retorno"],
            tam=12, cor=TEXT, espaco=1.35)
        txt(s, 9.3, 4.62, 3.4, 0.3, "ingredientes julgados", tam=8, cor=ORANGE,
            maiuscula=True)
        txt(s, 9.3, 4.98, 3.41, 1.5, [
            "estabilidade      ENTRA",
            "coerência          ENTRA",
            "volume               VETO",
            "proximidade     REPROVADO",
        ], tam=10, cor=TEXT, espaco=1.5)
        txt(s, 9.3, 6.30, 3.41, 0.3, "185 de 2.795 views vetadas", tam=9,
            cor=MUTED)
        rodape(s, "o número das views vetadas fica no slide, não na fala")
    return s_build(prs, "04", "3 de 4 — de onde vem o Ω", "0:35", "omega",
                   detalhe)


def s5_tatica(prs):
    def detalhe(s):
        txt(s, L, 4.62, 6.0, 0.3, "a camada tática", tam=8, cor=ORANGE,
            maiuscula=True)
        linhas_tabela(s, L, 4.95, 8.37, ["mercado que lê", "janela",
                                         "livro · neutro ao índice"], [
            ["recessão nos EUA", "3 · 5 · 10 pregões", "XLP defensivo × XLK cíclico"],
            ["maioria na Câmara", "20 pregões", "XLP defensivo × XLK cíclico"],
        ], tam=10, alt=0.34, cols=[0.31, 0.27, 0.42])
        txt(s, L, 6.02, 8.37, 0.6, ["tamanho pela mesma conta que dimensiona uma view",
                                    "zero parâmetro novo"],
            tam=11, cor=TEXT, espaco=1.35)
        caixa(s, 9.3, 4.90, 3.41, 1.35, cor_linha=ORANGE, espessura=1.5)
        txt(s, 9.5, 5.05, 3.0, 0.3, "o teto é um só", tam=11, cor=ORANGE,
            negrito=True)
        txt(s, 9.5, 5.38, 3.0, 0.7, ["Ω diz quanta confiança",
                                     "o teto diz quanto espaço"], tam=10,
            cor=TEXT, espaco=1.4)
        txt(s, 9.3, 6.35, 3.41, 0.3, "entrar não é somar, é dividir", tam=12,
            cor=TEXT, negrito=True)
        rodape(s, "estado final do diagrama — fica na tela enquanto o arco fecha")
    return s_build(prs, "05", "4 de 4 — camada tática e o teto", "0:35",
                   "tatica", detalhe)


def s6_ia(prs):
    s = slide_novo(prs)
    cabecalho(s, "06", "a infraestrutura de IA", "0:40")
    placeholder(s, L, 1.30, 6.10, 3.65, "print",
                "árvore de arquivos do second-brain",
                ["LOG.md · Decisoes_pendentes.md · Dump/analises/",
                 "captura real do repositório, sem edição"])
    txt(s, 7.05, 1.30, 5.66, 0.3, "as regras, escritas antes", tam=8,
        cor=ORANGE, maiuscula=True)
    txt(s, 7.05, 1.66, 5.66, 1.5, [
        "um dono por módulo · um branch por dono",
        "ritual obrigatório de abertura e fechamento",
        "toda pesquisa e decisão vira registro",
    ], tam=13, cor=TEXT, espaco=1.55)
    caixa(s, 7.05, 3.30, 5.66, 0.62, cor_linha=ORANGE, espessura=1.5)
    txt(s, 7.22, 3.48, 5.3, 0.3, "a IA nunca fecha decisão metodológica",
        tam=12, cor=ORANGE, negrito=True)
    txt(s, 7.05, 4.20, 5.66, 0.3, "o que o registro mede", tam=8, cor=ORANGE,
        maiuscula=True)
    for i, (n, r) in enumerate([("91", "sessões"), ("81", "decisões"),
                                ("232k", "tokens/sessão"), ("37%", "2ª rodada")]):
        x = 7.05 + i * 1.42
        txt(s, x, 4.55, 1.35, 0.4, n, tam=20, cor=TEXT, negrito=True)
        txt(s, x, 5.00, 1.35, 0.3, r, tam=8.5, cor=MUTED)
    txt(s, L, 5.25, 6.10, 0.9, [
        "nenhuma sessão recomeçava do zero",
        "o conhecimento parou de morar",
        "na cabeça de quem estava no teclado",
    ], tam=12, cor=TEXT, espaco=1.4)
    rodape(s, "IA no meio da apresentação, não no fim — o guia avisa que fechar "
              "em IA costuma ser a escolha errada")
    return s


def s7_pesquisa(prs):
    s = slide_novo(prs)
    cabecalho(s, "07", "o que a pesquisa descobriu sobre o sinal", "0:40")
    txt(s, L, 1.28, 12.09, 0.6,
        "O Polymarket reprecifica de uma vez, não em rampa.", tam=26, cor=TEXT,
        negrito=True)
    regua(s, L, 2.05, R)
    for i, (tit, itens, cor_t) in enumerate([
        ("ONDE NÃO HÁ SINAL",
         ["movimento  →  índice", "variance ratio ≈ 1", "autocorrelação ≈ 0",
          "", "→ por isso as views leem o NÍVEL"], MUTED),
        ("ONDE HÁ SINAL",
         ["movimento  →  relativo entre setores", "o único t-stat de verdade",
          "", "", "→ por isso a tática é um SPREAD"], ORANGE),
    ]):
        x = L + i * 6.16
        caixa(s, x, 2.30, 5.93, 2.15, cor_linha=RULE if i == 0 else ORANGE,
              espessura=0.75 if i == 0 else 1.5)
        txt(s, x + 0.22, 2.48, 5.5, 0.3, tit, tam=9, cor=cor_t, negrito=True)
        txt(s, x + 0.22, 2.85, 5.5, 1.4, itens, tam=12, cor=TEXT, espaco=1.35)
    placeholder(s, L, 4.75, 7.55, 1.95, "gráfico",
                "funil: 13 candidatas medidas  →  1 entrou",
                ["nomes das reprovadas legíveis — a densidade é o argumento"])
    txt(s, 8.55, 4.75, 4.16, 0.3, "a régua de entrada", tam=8, cor=ORANGE,
        maiuscula=True)
    txt(s, 8.55, 5.12, 4.16, 1.5, [
        "entra por acertar a probabilidade",
        "nunca por render mais",
        "",
        "momentum e velocidade de ajuste:",
        "passeio aleatório quando medidos",
    ], tam=11, cor=TEXT, espaco=1.4)
    rodape(s, "a conclusão vem primeiro; o funil entra como evidência, não como "
              "troféu")
    return s


def s8_resultados(prs):
    s = slide_novo(prs)
    cabecalho(s, "08", "resultados e crítica", "1:10")
    # as 3 anotacoes vao para as notas: sao instrucao de montagem, e impressas
    # no slide elas sozinhas ja empurram o 8 de volta para "slide de leitura"
    placeholder(s, L, 1.28, 7.55, 3.35, "gráfico",
                "curva acumulada  ×  SPY, com drawdown",
                ["3 anotações sobre a curva — ver notas"])
    placeholder(s, L, 4.80, 4.30, 1.90, "gráfico",
                "cascata da atribuição",
                ["mercado +28,7 · views +5,2 · custo −3,0",
                 "falta a 4ª barra: composição +2,2"])
    caixa(s, 5.05, 4.80, 3.12, 1.90, cor_linha=RULE)
    txt(s, 5.25, 4.98, 2.8, 0.24, "robustez", tam=8, cor=ORANGE, maiuscula=True)
    txt(s, 5.25, 5.30, 2.8, 1.2, [
        "4 botões girados",
        "+0,5 a +15,8 pp",
        "nenhuma combinação negativa",
    ], tam=11, cor=TEXT, espaco=1.45)
    txt(s, 8.55, 1.28, 4.16, 0.3, "os três números", tam=8, cor=ORANGE,
        maiuscula=True)
    txt(s, 8.55, 1.62, 4.16, 1.0, ["+33,2%  ×  +30,1%"], tam=22, cor=TEXT,
        negrito=True)
    txt(s, 8.55, 2.12, 4.16, 0.3, "ao mesmo risco — vol 17,6% × 17,9%", tam=10,
        cor=MUTED)
    linhas_tabela(s, 8.55, 2.60, 4.16, ["métrica", "kairos", "spy"], [
        ["Sharpe", "1,18", "1,08"],
        ["Máx. drawdown", "−19,6%", "−18,8%"],
        ["Alpha anual.", "+2,57%", "—"],
        ["Beta", "0,95", "1,00"],
    ], tam=9, alt=0.28, cols=[0.44, 0.28, 0.28])
    txt(s, 8.55, 4.35, 4.16, 0.3, "a crítica — falada por cima dos gráficos",
        tam=8, cor=ORANGE, maiuscula=True)
    txt(s, 8.55, 4.72, 4.16, 1.8, [
        "veio da subida, não da defesa",
        "uma janela, um regime",
        "ganha sozinha, perde somada",
        "42% do giro volta atrás",
    ], tam=11.5, cor=TEXT, espaco=1.5)
    txt(s, 8.55, 6.20, 4.16, 0.3, "+42,6 pp isolada  ×  −2,94 na carteira",
        tam=9, cor=MUTED)
    txt(s, L, 6.40, 7.55, 0.3,
        "a seguir: teto móvel · banda de não-negociação · testar em queda",
        tam=9, cor=MUTED)
    rodape(s, "\n".join([
        "Cada crítica é dita sobre o gráfico que a sustenta — nada de ler os",
        "quatro pontos com a tela parada.",
        "",
        "Anotações a desenhar sobre a curva:",
        "· na cava de abr/2025 — caiu junto, e mais fundo",
        "· no trecho de alta — é daqui que vem a vantagem",
        "· no eixo, ao final — 374 pregões, um único regime",
    ]))
    return s


def s9_fecho(prs):
    s = slide_novo(prs)
    cabecalho(s, "09", "fecho", "0:15")
    txt(s, L, 2.35, 12.09, 1.8, [
        "Q e Ω deixam de ser declarados.",
        "Passam a ser medidos.",
    ], tam=40, cor=TEXT, negrito=True, espaco=1.25)
    regua(s, L, 4.60, 7.4)
    txt(s, L, 4.85, 9.0, 0.8, [
        "o resultado depende de acertar a probabilidade —",
        "não de escolher bem o parâmetro",
    ], tam=15, cor=ORANGE, espaco=1.35)
    txt(s, L, 6.10, 9.0, 0.3, "o próximo passo é ver até onde essa ideia vai",
        tam=12, cor=MUTED)
    rodape(s, "cartão de texto: a tela esvazia de propósito no último parágrafo")
    return s


# ------------------------------------------------------------------- montagem
def monta(saida: Path = SAIDA) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for construtor in (s1_capa, s2_bl, s3_q, s4_omega, s5_tatica, s6_ia,
                       s7_pesquisa, s8_resultados, s9_fecho):
        construtor(prs)
    saida.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(saida))
    return saida


# ------------------------------------------------------------------ auto-teste
def _texto_dos_slides(prs):
    for i, slide in enumerate(prs.slides, start=1):
        linhas = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                linha = "".join(r.text for r in p.runs).strip()
                if linha:
                    linhas.append(linha)
        yield i, linhas


def demo() -> None:
    """Checa o que o .pptx nao denuncia: slide virando parágrafo e caixa fora
    da moldura. Roda com `python scripts/draft_semis_pptx.py --demo`."""
    from pptx import Presentation as _P

    caminho = monta()
    prs = _P(str(caminho))
    lados = (prs.slide_width, prs.slide_height)

    for shape in (sh for s in prs.slides for sh in s.shapes):
        assert shape.left >= -Emu(1), f"caixa à esquerda da moldura: {shape.shape_id}"
        assert shape.top >= -Emu(1), f"caixa acima da moldura: {shape.shape_id}"
        assert shape.left + shape.width <= lados[0] + Emu(9144), \
            f"caixa estoura a direita: {shape.shape_id}"
        assert shape.top + shape.height <= lados[1] + Emu(9144), \
            f"caixa estoura embaixo: {shape.shape_id}"

    # As notas de producao ficam no painel de anotacoes, entao nada abaixo conta
    # texto que a banca nao ve. O criterio que importa nao e o total de palavras
    # (um diagrama tem muitos rotulos curtos e continua limpo) e sim quantas
    # LINHAS parecem frase: e isso que transforma slide de apoio em slide de
    # leitura.
    pico_frases = 0
    for numero, linhas in _texto_dos_slides(prs):
        palavras = sum(len(l.split()) for l in linhas)
        assert palavras <= MAX_PALAVRAS, \
            f"slide {numero} com {palavras} palavras (teto {MAX_PALAVRAS})"
        for linha in linhas:
            assert len(linha.split()) <= MAX_PALAVRAS_LINHA, \
                f"slide {numero}: linha virou frase corrida — {linha!r}"
            assert len(linha) <= MAX_LINHA, \
                f"slide {numero}: linha longa demais — {linha!r}"
        frases = [l for l in linhas if len(l.split()) >= 8]
        pico_frases = max(pico_frases, len(frases))
        assert len(frases) <= MAX_FRASES, (
            f"slide {numero} tem {len(frases)} linhas de frase (teto "
            f"{MAX_FRASES}) — é slide para falar por cima, não para ler")

    assert len(prs.slides._sldIdLst) == 9, "o deck tem 9 slides"
    print(f"ok — {caminho.name}: 9 slides · no pior slide, {pico_frases} linhas "
          f"de frase (teto {MAX_FRASES})")


if __name__ == "__main__":
    import sys

    if "--demo" in sys.argv:
        demo()
    else:
        print(monta())

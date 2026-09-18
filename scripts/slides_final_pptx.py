"""Slides novos da final — 3 (Polymarket), 4 (comparação) e 12 (second brain).

Gera `Final/Slides_novos.pptx` com três grupos, no estilo restilizado do deck da
semifinal (fundo radial, Bahnschrift + Consolas, cards em degradê, número
fantasma, halos de canto). Os slides entram no deck da final por copiar/colar
no PowerPoint com "manter formatação de origem".

  3  · como funciona o Polymarket: o card de um mercado real (o Fed de
       setembro/2025) e a probabilidade de cada desfecho andando até virar 0 ou 1.
  4  · o mesmo dia, duas previsões: para cada ETF, o retorno esperado no dia do
       anúncio segundo o MERCADO (π) e segundo o POLYMARKET (π + β·surpresa).
       Dois slides com Morph — véspera de 27/01/2026 (concordam) e véspera de
       16/09/2025 (divergem). Dado de `comparacao_etf_s4.py`.
  12 · o second brain acendendo: o grafo do repositório (`grafo_repo.py`) e o
       caminho REAL que um prompt percorreu, medido no transcript da sessão por
       `trace_sessao.py`. Uma sequência de slides Morph com avanço automático:
       um clique dispara o caminho inteiro.
  13 · IA em números: dois gráficos da página 5 do relatório
       (`Uteis/graficos/p5_*.png`, de `graficos_p5.py`) — erros × quem pegou e
       os donuts do esforço por tipo — e os quatro indicadores na faixa de
       baixo. Texto só nos kickers: o slide é a base de uma leitura geral sobre
       IA em estratégia quant, não dos nossos números.
  16 · o estudo estatístico do Polymarket (`estudo_polymarket.py`,
       `Final/estudo/Estudo_polymarket.md`) em duas observações e uma conclusão:
       calibração (0,043) e decisão do Fed (0,9 × 4,8 bps). Três slides de
       apêndice para a arguição — A1 acurácia por horizonte; A2 Polymarket ×
       futuro de FF e lead-lag; A3 event-study e o placar "onde está o valor".
  14 · o slide 28 do deck da final (placar 2 × 2, cascata, curva e drawdown)
       com UM tile a mais: P(excesso > 0) e o IC95 do bootstrap
       (`analise_backtest.py`) — sorte ou habilidade?
  15 · análise crítica em três números com leitura qualitativa: 457 pregões
       de track record mínimo (ainda não provada) · +6,2 bps por dia de queda
       do SPY (defensiva) · 3 de 7 no protocolo de Arnott, Harvey e Markowitz.
       A9 · tear-sheet completa, rolling IR e a curva do DSR.

Como o Morph casa as formas: nome com prefixo `!!` = mesmo objeto entre slides
(anda, muda de cor, muda de tamanho); todo o resto ganha nome único do slide
(`sela`) para o PowerPoint não casar duas caixas sem relação. Requer PowerPoint
2016+/365; em outros leitores a transição cai no `<mc:Fallback>` e vira fade.

Uso:

    python scripts/slides_final_pptx.py             # gera o deck DO ZERO (apaga edições manuais)
    python scripts/slides_final_pptx.py --anexar    # só acrescenta 14/15/A9 ao deck salvo
    python scripts/slides_final_pptx.py --transicao # insere a transição "Resposta da hipótese" após o 15
    python scripts/slides_final_pptx.py --render    # e exporta PNG via PowerPoint
    python scripts/slides_final_pptx.py --demo      # auto-teste
"""

import json
import sys
from pathlib import Path

import matplotlib
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Emu, Inches, Pt

matplotlib.use("Agg")
import matplotlib.dates as mdates          # noqa: E402
import matplotlib.pyplot as plt            # noqa: E402

import grafo_repo                          # noqa: E402
import graficos_p5                         # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
SAIDA = RAIZ / "Final" / "Slides_novos.pptx"
GRAFICOS = RAIZ / "Uteis" / "graficos"
DADOS = RAIZ / "Uteis" / "dados"
FED_PARQUET = RAIZ / "data" / "polymarket_fed_reunioes.parquet"

# ------------------------------------------------------------------ paleta
# Lida do XML do deck restilizado (Semis/KAIROS_semifinal.pptx, sessão 44).
ORANGE = RGBColor(0xFF, 0xB5, 0x31)
BLUE = RGBColor(0x4F, 0x8E, 0xF7)
MINT = RGBColor(0x2F, 0xD3, 0xB0)
TEXT = RGBColor(0xF5, 0xF7, 0xFB)
BODY = RGBColor(0xBA, 0xC6, 0xDA)
MUTED = RGBColor(0x83, 0x95, 0xB5)
FOOT = RGBColor(0x58, 0x69, 0x8A)
RULE = RGBColor(0x2A, 0x3F, 0x63)
GHOST = RGBColor(0x10, 0x1A, 0x2E)
CARD_A, CARD_B = "182946", "0F182B"       # degradê do card
BORDA_A, BORDA_B = "35507A", "22345A"     # degradê da borda

COND = "Bahnschrift SemiBold Condensed"
SEMI = "Bahnschrift SemiBold"
LIGHT = "Bahnschrift SemiLight"
BAHN = "Bahnschrift"
MONO = "Consolas"
SERIF = "Georgia"

L, R = 0.80, 12.53                        # margens do deck restilizado
Y_TITULO = 0.93
PE = 7.167                                # régua do rodapé

FAMILIA_COR = {"contrato": ORANGE, "analise": MINT, "codigo": BLUE,
               "entrega": MUTED}
FAMILIA_ROT = {"contrato": "contrato", "analise": "análise", "codigo": "código",
               "entrega": "entrega"}

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


# --------------------------------------------------------------- primitivas
def hexa(cor):
    return f"{cor[0]:02X}{cor[1]:02X}{cor[2]:02X}"


def alpha(sh, val):
    """Transparência num preenchimento sólido (val em milésimos de %)."""
    sh.fill.fore_color._xFill.find(A + "srgbClr").append(
        parse_xml(f'<a:alpha {nsdecls("a")} val="{val}"/>'))


def efeito(sh, xml):
    sh.shadow.inherit = False
    sp = sh._element.spPr
    velho = sp.find(A + "effectLst")
    if velho is not None:
        sp.remove(velho)
    sp.append(parse_xml(f'<a:effectLst {nsdecls("a")}>{xml}</a:effectLst>'))


def brilho(cor, rad=90000, alfa=40000):
    return (f'<a:glow rad="{rad}"><a:srgbClr val="{hexa(cor)}">'
            f'<a:alpha val="{alfa}"/></a:srgbClr></a:glow>')


def caixa(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return sh


def paragrafo(tf, runs, alinha=PP_ALIGN.LEFT, primeiro=False, entre=1.0,
              depois=0.0, antes=0.0):
    """Um parágrafo; `runs` = [(texto, fonte, pt, cor, negrito, spc, itálico)]."""
    p = tf.paragraphs[0] if primeiro else tf.add_paragraph()
    p.alignment = alinha
    p.line_spacing = entre
    p.space_after = Pt(depois)
    p.space_before = Pt(antes)
    for texto, fonte, pt, cor, *resto in runs:
        negrito = resto[0] if len(resto) > 0 else False
        spc = resto[1] if len(resto) > 1 else 0
        italico = resto[2] if len(resto) > 2 else False
        r = p.add_run()
        r.text = texto
        f = r.font
        f.name, f.size, f.bold, f.italic = fonte, Pt(pt), negrito, italico
        f.color.rgb = cor
        if spc:
            r._r.get_or_add_rPr().set("spc", str(spc))
    return p


def texto(slide, x, y, w, h, runs, alinha=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          entre=1.0, wrap=True):
    sh = caixa(slide, x, y, w, h, anchor, wrap)
    paragrafo(sh.text_frame, runs, alinha, primeiro=True, entre=entre)
    return sh


def kicker(slide, x, y, w, rotulo, cor=ORANGE):
    return texto(slide, x, y, w, 0.2, [(rotulo, MONO, 9, cor, True, 180)])


def card(slide, x, y, w, h, raio=4000):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),
                                Inches(y), Inches(w), Inches(h))
    sp = sh._element.spPr
    for tag in ("solidFill", "ln", "prstGeom"):
        el = sp.find(A + tag)
        if el is not None:
            sp.remove(el)
    sp.append(parse_xml(
        f'<a:prstGeom {nsdecls("a")} prst="roundRect"><a:avLst>'
        f'<a:gd name="adj" fmla="val {raio}"/></a:avLst></a:prstGeom>'))
    sp.append(parse_xml(
        f'<a:gradFill {nsdecls("a")} rotWithShape="1"><a:gsLst>'
        f'<a:gs pos="0"><a:srgbClr val="{CARD_A}"/></a:gs>'
        f'<a:gs pos="100000"><a:srgbClr val="{CARD_B}"/></a:gs></a:gsLst>'
        '<a:lin ang="5400000" scaled="0"/></a:gradFill>'))
    sp.append(parse_xml(
        f'<a:ln {nsdecls("a")} w="9525"><a:gradFill rotWithShape="1"><a:gsLst>'
        f'<a:gs pos="0"><a:srgbClr val="{BORDA_A}"/></a:gs>'
        f'<a:gs pos="100000"><a:srgbClr val="{BORDA_B}"/></a:gs></a:gsLst>'
        '<a:lin ang="5400000" scaled="0"/></a:gradFill></a:ln>'))
    efeito(sh, "")
    sh.text_frame.text = ""
    return sh


def fio(slide, x1, y1, x2, y2, cor=RULE, pt=0.75, alfa=None):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),
                                    Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = cor
    ln.line.width = Pt(pt)
    if alfa is not None:
        ln.line.color._xFill.find(A + "srgbClr").append(
            parse_xml(f'<a:alpha {nsdecls("a")} val="{alfa}"/>'))
    return ln


def regua_dourada(slide, x1, y, x2):
    """A régua do cabeçalho: dourado que some para a direita."""
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),
                                    Inches(y), Inches(x2), Inches(y))
    sp = ln._element.spPr
    for el in sp.findall(A + "ln"):
        sp.remove(el)
    sp.append(parse_xml(
        f'<a:ln {nsdecls("a")} w="12700"><a:gradFill rotWithShape="1"><a:gsLst>'
        f'<a:gs pos="0"><a:srgbClr val="{hexa(ORANGE)}"/></a:gs>'
        f'<a:gs pos="30000"><a:srgbClr val="{hexa(RULE)}"/></a:gs>'
        f'<a:gs pos="100000"><a:srgbClr val="{hexa(RULE)}"><a:alpha val="30000"/>'
        '</a:srgbClr></a:gs></a:gsLst><a:lin ang="0" scaled="0"/></a:gradFill></a:ln>'))
    return ln


def ponto(slide, x, y, d, cor, luz=None, alfa=None):
    """Disco centrado em (x, y) com diâmetro d, opcionalmente aceso."""
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - d / 2),
                                Inches(y - d / 2), Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = cor
    if alfa is not None:
        alpha(sh, alfa)
    sh.line.fill.background()
    efeito(sh, brilho(cor, *luz) if luz else "")
    sh.text_frame.text = ""
    return sh


def halo_canto(slide, nome, x, y, cor, pico):
    """Os halos de canto do deck: elipse com degradê radial que some."""
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(x), Emu(y), Emu(7000000),
                                Emu(5200000))
    sh.name = nome
    sp = sh._element.spPr
    for tag in ("solidFill", "ln"):
        el = sp.find(A + tag)
        if el is not None:
            sp.remove(el)
    paradas = [(0, pico), (45000, pico * 4 // 10), (80000, 1000), (100000, 0)]
    gs = "".join(f'<a:gs pos="{p}"><a:srgbClr val="{hexa(cor)}"><a:alpha val="{a}"/>'
                 '</a:srgbClr></a:gs>' for p, a in paradas)
    sp.append(parse_xml(f'<a:gradFill {nsdecls("a")} rotWithShape="1"><a:gsLst>{gs}'
                        '</a:gsLst><a:path path="circle"><a:fillToRect l="50000" '
                        't="50000" r="50000" b="50000"/></a:path></a:gradFill>'))
    sp.append(parse_xml(f'<a:ln {nsdecls("a")}><a:noFill/></a:ln>'))
    sh.text_frame.text = ""
    return sh


def imagem(slide, caminho, x, y, w, h=None):
    return slide.shapes.add_picture(str(caminho), Inches(x), Inches(y), Inches(w),
                                    None if h is None else Inches(h))


# -------------------------------------------------------------- transições
def morph(slide, dur_ms=600, avanca_ms=None):
    """Morph AO ENTRAR neste slide; `avanca_ms` sai dele sozinho depois disso.

    O python-pptx não tem API de transição; o XML vai direto no `<p:sld>`, no
    lugar que o schema espera (depois de `clrMapOvr`). `advTm` é o avanço
    automático: é ele que faz a sequência do slide 12 tocar com um clique só.
    """
    adv = f' advTm="{avanca_ms}"' if avanca_ms is not None else ""
    slide.element.append(parse_xml(
        '<mc:AlternateContent'
        ' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
        ' xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        '<mc:Choice Requires="p14"'
        ' xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main">'
        f'<p:transition spd="slow"{adv} p14:dur="{dur_ms}">'
        '<p159:morph option="byObject"'
        ' xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main"/>'
        '</p:transition></mc:Choice>'
        f'<mc:Fallback><p:transition spd="fast"{adv}><p:fade/></p:transition>'
        '</mc:Fallback></mc:AlternateContent>'))


def sela(slide, tag):
    for i, sh in enumerate(slide.shapes):
        if not sh.name.startswith("!!"):
            sh.name = f"{tag}_{i}"


# ------------------------------------------------------------------ moldura
def moldura(prs, numero, secao, titulo):
    """Fundo radial, halos, número fantasma, cabeçalho, título e rodapé."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    cSld = s._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}cSld")
    cSld.insert(0, parse_xml(
        f'<p:bg {nsdecls("p", "a")}><p:bgPr><a:gradFill rotWithShape="1"><a:gsLst>'
        '<a:gs pos="0"><a:srgbClr val="10203A"/></a:gs>'
        '<a:gs pos="55000"><a:srgbClr val="0A1325"/></a:gs>'
        '<a:gs pos="100000"><a:srgbClr val="04070E"/></a:gs></a:gsLst>'
        '<a:path path="circle"><a:fillToRect l="38000" r="62000" b="100000"/></a:path>'
        '</a:gradFill><a:effectLst/></p:bgPr></p:bg>'))
    halo_canto(s, "!!glow_tl", -2600000, -2200000, BLUE, 10000)
    halo_canto(s, "!!glow_br", 8200000, 4600000, ORANGE, 7000)

    fantasma = texto(s, 7.546, -0.252, 5.413, 2.187, [(numero, COND, 150, GHOST)],
                     alinha=PP_ALIGN.RIGHT, wrap=False)
    fantasma.name = "!!ghost"
    tick = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(603504), Emu(446024),
                              Emu(45720), Emu(164592))
    tick.fill.solid()
    tick.fill.fore_color.rgb = ORANGE
    tick.line.fill.background()
    tick.name = "!!tick"

    texto(s, L, 0.467, 8.0, 0.222, [("KAIROS", MONO, 12, ORANGE, True, 300),
                                    (f"   {secao}", MONO, 11, MUTED, False, 72)],
          anchor=MSO_ANCHOR.MIDDLE)
    texto(s, 0, 0.511, 12.53, 0.167, [(numero, MONO, 9, MUTED, False, 45)],
          alinha=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    regua_dourada(s, L, 0.771, R)
    texto(s, L, Y_TITULO, R - L, 0.7, [(titulo, SEMI, 20, TEXT)])

    fio(s, L, PE, R, PE)
    pe = texto(s, L, 7.276, 6.56, 0.175,
               [("ITAÚ QUANT AI CHALLENGE  ·  FINAL 2026", MONO, 7, FOOT, False, 120)],
               wrap=False)
    pe.name = "!!foot_l"
    pd_ = texto(s, 5.968, 7.276, 6.56, 0.175,
                [("BLACK-LITTERMAN  ×  POLYMARKET", MONO, 7, FOOT, False, 120)],
                alinha=PP_ALIGN.RIGHT, wrap=False)
    pd_.name = "!!foot_r"
    return s


# =================================================================== slide 3
def serie_fed_set25():
    """Probabilidade de cada desfecho da reunião de 17/09/2025, pré-abertura."""
    d = pd.read_parquet(FED_PARQUET)
    d = d[d["mercado"].str.contains("September 2025")]
    w = d.pivot_table(index="data", columns="mercado", values="probabilidade")
    volume = d.groupby("mercado")["volume"].first().sum()
    curto = {}
    for c in w.columns:
        if "decreases" in c and "25 bps" in c:
            curto[c] = "corte de 25 bps"
        elif "decreases" in c:
            curto[c] = "corte de 50+ bps"
        elif "increases" in c:
            curto[c] = "alta de 25+ bps"
        else:
            curto[c] = "sem mudança"
    w = w.rename(columns=curto)
    return w[["corte de 25 bps", "sem mudança", "corte de 50+ bps", "alta de 25+ bps"]], volume


CORES_FED = {"corte de 25 bps": ORANGE, "sem mudança": BLUE,
             "corte de 50+ bps": MINT, "alta de 25+ bps": MUTED}
MARCOS = [("2025-07-30", "30/07 · Fed mantém  ", "right"),
          ("2025-08-01", "  01/08 · payroll fraco", "left"),
          ("2025-09-17", "17/09 · corta 25 bps  ", "right")]


def grafico_fed(w, destino):
    """As quatro linhas andando até 0 ou 1 — fundo transparente, tipos do deck."""
    plt.rcParams["font.family"] = [BAHN, "Segoe UI"]
    fig, ax = plt.subplots(figsize=(7.0, 3.55))
    w = w.loc["2025-07-14":"2025-09-17"]
    fim = {}
    for nome, cor in CORES_FED.items():
        s = w[nome].dropna()
        ax.plot(s.index, s.values * 100, color="#" + hexa(cor), lw=2.0 if nome.startswith("corte de 25") else 1.4,
                solid_capstyle="round", zorder=3)
        fim[nome] = (s.index[-1], s.values[-1] * 100)
    # rótulos no fim das linhas, empurrados para não se sobrepor (7 pp de vão)
    y_rot, ultimo = {}, -99.0
    for nome in sorted(fim, key=lambda n: fim[n][1]):
        ultimo = max(fim[nome][1], ultimo + 7.0)
        y_rot[nome] = ultimo
    for nome, cor in CORES_FED.items():
        ax.annotate(nome, fim[nome], xytext=(6, y_rot[nome] - fim[nome][1]),
                    textcoords="offset points", va="center", fontsize=8.5,
                    color="#" + hexa(cor))
    for dia, rot, lado in MARCOS:
        x = pd.Timestamp(dia)
        ax.axvline(x, color="#" + hexa(RULE), lw=0.8, ls=(0, (3, 3)), zorder=1)
        ax.text(x, 104, rot, fontsize=7.5, color="#" + hexa(MUTED), ha=lado,
                va="bottom")
    ax.set_ylim(0, 100)
    ax.set_xlim(w.index.min(), w.index.max() + pd.Timedelta(days=1))
    ax.set_yticks([0, 25, 50, 75, 100])
    ax.set_yticklabels(["0", "25", "50", "75", "100%"], fontsize=8, color="#" + hexa(MUTED))
    ax.xaxis.set_major_locator(mdates.MonthLocator(bymonthday=[1, 15]))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%d/%m"))
    ax.tick_params(axis="x", labelsize=8, colors="#" + hexa(MUTED), length=0)
    ax.tick_params(axis="y", length=0)
    for lado in ("top", "right", "left"):
        ax.spines[lado].set_visible(False)
    ax.spines["bottom"].set_color("#" + hexa(RULE))
    ax.grid(axis="y", color="#" + hexa(RULE), lw=0.5, alpha=0.6)
    ax.set_ylabel("preço do contrato \"Sim\"  (¢ = %)", fontsize=8, color="#" + hexa(MUTED))
    fig.subplots_adjust(0.07, 0.10, 0.86, 0.88)
    destino.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(destino, transparent=True, dpi=200)
    plt.close(fig)
    return destino


def s3_polymarket(prs, numero="03"):
    w, volume = serie_fed_set25()
    png = grafico_fed(w, GRAFICOS / "s3_fed_set25.png")
    vespera = w.ffill().loc[:"2025-09-16 12:00"].iloc[-1]

    s = moldura(prs, numero, "Como funciona o Polymarket",
                "Um contrato paga US$ 1 se o evento acontece. O preço é a probabilidade.")

    # --- o "print" do mercado, com os números reais da véspera
    x, y, wc, hc = L, 1.75, 4.55, 3.25
    card(s, x, y, wc, hc)
    kicker(s, x + 0.28, y + 0.22, 4.0, "MERCADO REAL  ·  POLYMARKET")
    texto(s, x + 0.28, y + 0.50, 4.0, 0.4, [("Fed decision in September 2025?", SEMI, 14, TEXT)])
    fio(s, x + 0.28, y + 0.98, x + wc - 0.28, y + 0.98)
    ly = y + 1.10
    for nome in vespera.sort_values(ascending=False).index:
        p = float(vespera[nome])
        cor = ORANGE if nome.startswith("corte de 25") else MUTED
        texto(s, x + 0.28, ly, 2.2, 0.26, [(nome, LIGHT, 11.5, BODY if cor is MUTED else TEXT)],
              anchor=MSO_ANCHOR.MIDDLE)
        texto(s, x + wc - 1.05, ly, 0.77, 0.26, [(f"{round(p * 100):d}¢", MONO, 12, cor, True)],
              alinha=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        barra = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.28),
                                   Inches(ly + 0.30), Inches(max(0.03, 3.99 * p)), Inches(0.05))
        barra.fill.solid()
        barra.fill.fore_color.rgb = cor
        barra.line.fill.background()
        efeito(barra, brilho(ORANGE, 40000, 35000) if cor is ORANGE else "")
        ly += 0.46
    milhoes = f"{volume / 1e6:,.0f}".replace(",", ".")
    texto(s, x + 0.28, y + hc - 0.42, 4.0, 0.2,
          [(f"US$ {milhoes} mi negociados  ·  leitura de 16/09/2025, véspera",
            MONO, 8, MUTED)])

    # --- a probabilidade andando
    kicker(s, 5.75, 1.75, 6.5, "A PROBABILIDADE ANDA ATÉ VIRAR 0 OU 100")
    imagem(s, png, 5.55, 1.98, 6.98)

    # --- faixa de baixo: os dois números do slide 2 e como o Kairós lê
    yb = 5.28
    for i, (num, rot) in enumerate([("90,1%", "de acerto a 1 mês da resolução"),
                                    ("US$ 26,2 bi", "negociados no 1º tri de 2026")]):
        cx = L + i * 2.95
        card(s, cx, yb, 2.72, 1.6)
        texto(s, cx + 0.25, yb + 0.22, 2.3, 0.65, [(num, COND, 30, ORANGE)])
        texto(s, cx + 0.25, yb + 0.95, 2.3, 0.5, [(rot, BAHN, 9.5, BODY)], entre=1.05)
    cx = L + 2 * 2.95
    card(s, cx, yb, R - cx, 1.6)
    kicker(s, cx + 0.28, yb + 0.22, 5.0, "O QUE O KAIROS LÊ")
    tf = caixa(s, cx + 0.28, yb + 0.52, R - cx - 0.56, 1.0).text_frame
    paragrafo(tf, [("O preço de cada desfecho, duas vezes por dia. ", SEMI, 11.5, TEXT),
                   ("Sempre a leitura anterior à abertura da bolsa — nunca a do dia.", LIGHT, 11.5, BODY)],
              primeiro=True, entre=1.12, depois=4)
    paragrafo(tf, [("Quem errar perde dinheiro de verdade. ", SEMI, 11.5, TEXT),
                   ("É o que faz do preço uma opinião com custo, não uma pesquisa.", LIGHT, 11.5, BODY)],
              entre=1.12)
    sela(s, "s3")
    return s


# =================================================================== slide 4
ETFS = [("XLK", "tecnologia"), ("SPY", "S&P 500"), ("XLF", "financeiro"),
        ("XLE", "energia"), ("XLV", "saúde"), ("XLP", "consumo básico"),
        ("TIP", "Treasury · inflação"), ("XLU", "utilities"), ("TLT", "Treasury 20+")]
DIAS_S4 = [("2026-01-27", "concordam"), ("2025-09-16", "divergem")]
EIXO_MIN, EIXO_MAX = -1.8, 0.6            # % — o alcance dos dois dias
X0, X1 = 3.55, 11.95                      # onde o eixo vive no slide
Y0, PASSO = 2.72, 0.42                    # primeira linha e espaçamento


def x_de(valor_pct):
    return X0 + (valor_pct - EIXO_MIN) / (EIXO_MAX - EIXO_MIN) * (X1 - X0)


def fmt_pct(v):
    return f"{v * 100:+.2f}%".replace(".", ",").replace("+", "+ ").replace("-", "− ")


def s4_comparacao(prs, numero="04"):
    d = pd.read_csv(DADOS / "s4_comparacao.csv", index_col=0, parse_dates=True)
    slides = []
    for k, (dia, estado) in enumerate(DIAS_S4):
        r = d.loc[dia]
        if estado == "concordam":
            titulo = "Na maioria dos dias, Polymarket e mercado esperam a mesma coisa."
            lead, resto = ("Surpresa de 0 bps.  ",
                           "Polymarket e juros precificam a mesma decisão — o Kairós não tem opinião e carrega a carteira do mercado.")
        else:
            titulo = "Em alguns, divergem — e é aí que o Kairós age."
            lead, resto = (f"Surpresa de {r['surpresa_liquida']:+.0f} bps.  ",
                           "O Polymarket vê menos corte do que os juros embutem: dinheiro em risco discorda do preço. A conta vira retorno esperado por ETF.")
        s = moldura(prs, numero, "O mesmo dia, duas previsões", titulo)

        # faixa do dia — o card persiste, o texto troca (fade)
        cd = card(s, L, 1.72, R - L, 0.72)
        cd.name = "!!faixa"
        data_br = pd.Timestamp(dia).strftime("%d/%m/%Y")
        texto(s, L + 0.28, 1.86, 3.0, 0.45,
              [("VÉSPERA DO FED", MONO, 9, ORANGE, True, 180),
               (f"   {data_br}", MONO, 9, MUTED, False, 60)], anchor=MSO_ANCHOR.MIDDLE)
        texto(s, L + 3.3, 1.79, R - L - 3.6, 0.6,
              [(lead, SEMI, 11.5, TEXT), (resto, LIGHT, 11.5, BODY)],
              anchor=MSO_ANCHOR.MIDDLE, entre=1.1)

        # eixo
        y_top, y_bot = Y0 - 0.28, Y0 + PASSO * (len(ETFS) - 1) + 0.28
        for v in (-1.5, -1.0, -0.5, 0.0, 0.5):
            x = x_de(v)
            ln = fio(s, x, y_top, x, y_bot, RULE, 1.25 if v == 0 else 0.6,
                     None if v == 0 else 45000)
            ln.name = f"!!grade_{v}"
            t = texto(s, x - 0.5, y_bot + 0.04, 1.0, 0.2,
                      [(f"{v:+.1f}%".replace(".", ",").replace("+0,0", "0,0"), MONO, 8, MUTED)],
                      alinha=PP_ALIGN.CENTER)
            t.name = f"!!tick_{v}"
        t = texto(s, L, y_bot + 0.04, X0 - L - 0.2, 0.2,
                  [("retorno esperado no dia do anúncio  →", MONO, 8, MUTED, False, 60)],
                  wrap=False)
        t.name = "!!eixo_rot"

        # as nove linhas
        for i, (etf, setor) in enumerate(ETFS):
            y = Y0 + i * PASSO
            t = texto(s, L, y - 0.15, 0.7, 0.3, [(etf, MONO, 11.5, TEXT, True)],
                      anchor=MSO_ANCHOR.MIDDLE)
            t.name = f"!!etf_{etf}"
            t = texto(s, L + 0.72, y - 0.15, 2.0, 0.3, [(setor, LIGHT, 9.5, MUTED)],
                      anchor=MSO_ANCHOR.MIDDLE)
            t.name = f"!!setor_{etf}"
            ln = fio(s, X0, y, X1, y, RULE, 0.5, 35000)
            ln.name = f"!!trilho_{etf}"

            pi, poly = float(r[f"pi_{etf}"]), float(r[f"poly_{etf}"])
            xm, xp = x_de(pi * 100), x_de(poly * 100)
            lig = fio(s, xm, y, xp, y, ORANGE, 2.0, 55000)
            lig.name = f"!!lig_{etf}"
            pp = ponto(s, xp, y, 0.19, ORANGE, luz=(70000, 40000))
            pp.name = f"!!pm_{etf}"
            pm = ponto(s, xm, y, 0.13, BLUE, luz=(50000, 30000))
            pm.name = f"!!mk_{etf}"
            esquerda = poly < -0.0015
            t = texto(s, xp - 0.95 if esquerda else xp + 0.17, y - 0.13, 0.8, 0.26,
                      [(fmt_pct(poly), MONO, 9, ORANGE, True)],
                      alinha=PP_ALIGN.RIGHT if esquerda else PP_ALIGN.LEFT,
                      anchor=MSO_ANCHOR.MIDDLE)
            t.name = f"!!val_{etf}"

        # legenda
        yl = 6.66
        ponto(s, X0 + 0.08, yl + 0.1, 0.13, BLUE).name = "!!leg_mk"
        t = texto(s, X0 + 0.24, yl, 3.6, 0.2,
                  [("MERCADO", MONO, 8.5, TEXT, True, 100),
                   ("  o equilíbrio π — o prêmio normal do dia", LIGHT, 9, MUTED)],
                  anchor=MSO_ANCHOR.MIDDLE)
        t.name = "!!leg_mk_t"
        ponto(s, X0 + 4.05, yl + 0.1, 0.15, ORANGE).name = "!!leg_pm"
        t = texto(s, X0 + 4.23, yl, 4.6, 0.2,
                  [("POLYMARKET", MONO, 8.5, TEXT, True, 100),
                   ("  π + β × surpresa — β medido em reuniões passadas", LIGHT, 9, MUTED)],
                  anchor=MSO_ANCHOR.MIDDLE)
        t.name = "!!leg_pm_t"
        t = texto(s, L, 6.93, R - L, 0.2,
                  [("surpresa = o corte que o Polymarket precifica − o que o T-bill de 3 meses embute, "
                    "líquida do viés do instrumento  ·  apêndices A5 e A6", MONO, 7.5, FOOT)])
        t.name = "!!nota"

        # a régua e a nota de rodapé ficam; anotar o dado bruto para a banca
        s.notes_slide.notes_text_frame.text = (
            f"{dia}: E_poly = {r['e_poly_bps']:+.1f} bps · juros (DTB3 − DFF) = {r['e_ff_bps']:+.1f} bps · "
            f"surpresa bruta {r['surpresa_bps']:+.1f} · líquida {r['surpresa_liquida']:+.1f} bps · "
            f"{int(r['dias_ate_evento'])} pregão até a reunião.")
        if k > 0:
            morph(s, 900)
        sela(s, f"s4{'ab'[k]}")
        slides.append(s)
    return slides


# ================================================================== slide 12
GX, GY, GW = 0.62, 1.55, 5.45             # onde o grafo entra no slide
CX = 6.45                                 # a coluna da direita
PASSO_MS = 320                            # dwell entre passos; o Morph dura 550


def grafo_para_slide():
    """(png, posições por nome em polegadas do slide, nº nós, nº arestas)."""
    nomes, familias, arestas = grafo_repo.levantar()
    pos = grafo_repo.layout(len(nomes), arestas)
    ext = float(np.abs(pos).max()) * 1.08
    png = grafo_repo.desenhar(nomes, familias, arestas, pos,
                              GRAFICOS / "s12_grafo.png", rotulos=3, miudos=0,
                              extensao=ext)
    # o PNG é quadrado, eixos em [-ext, ext] e sem recorte: mapa linear
    xy = {n: (GX + (x + ext) / (2 * ext) * GW, GY + (ext - y) / (2 * ext) * GW)
          for n, (x, y) in zip(nomes, pos)}
    return png, xy, len(nomes), len(arestas)


def s12_second_brain(prs, numero="12"):
    trace = json.loads((DADOS / "trace_sessao.json").read_text(encoding="utf-8"))
    png, xy, n_nos, n_arestas = grafo_para_slide()
    arquivos = trace["arquivos"]
    faltando = [a["nome"] for a in arquivos if a["nome"] not in xy]
    assert not faltando, f"arquivos do trace fora do grafo: {faltando}"
    passos = len(arquivos)
    slides = []
    for k in range(passos + 2):           # 0 = repouso · 1..n = acende · n+1 = pousa
        final = k == passos + 1
        acesos = arquivos[:min(k, passos)]
        s = moldura(prs, numero, "Setup de IA",
                    "Um prompt não lê o repositório inteiro: acende um caminho.")

        img = imagem(s, png, GX, GY, GW, GW)
        img.name = "!!grafo"
        veu = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(GX), Inches(GY),
                                 Inches(GW), Inches(GW))
        veu.fill.solid()
        veu.fill.fore_color.rgb = RGBColor(0x04, 0x07, 0x0E)
        alpha(veu, 42000 if acesos else 0)   # escurece quando o caminho acende
        veu.line.fill.background()
        efeito(veu, '<a:softEdge rad="700000"/>')   # sem borda dura no fundo
        veu.text_frame.text = ""
        veu.name = "!!veu"
        t = texto(s, GX + 0.1, GY + GW - 0.02, GW, 0.2,
                  [(f"{n_nos} arquivos  ·  {n_arestas} ligações  ·  medido no branch por grafo_repo.py",
                    MONO, 7.5, FOOT)], alinha=PP_ALIGN.CENTER)
        t.name = "!!grafo_leg"

        # o caminho: fio entre nós consecutivos, nó aceso, e o pulso no atual
        for i, a in enumerate(acesos):
            x, y = xy[a["nome"]]
            cor = FAMILIA_COR[a["familia"]]
            if i > 0:
                x0, y0 = xy[acesos[i - 1]["nome"]]
                ln = fio(s, x0, y0, x, y, ORANGE, 1.5, 70000)
                ln.name = f"!!e_{i}"
            p = ponto(s, x, y, 0.2, cor, luz=(100000, 60000))
            p.name = f"!!n_{i}"
        if acesos and not final:
            x, y = xy[acesos[-1]["nome"]]
            pulso = ponto(s, x, y, 0.55, ORANGE, alfa=40000)
            efeito(pulso, '<a:softEdge rad="150000"/>')
            pulso.name = "!!pulso"

        # --- coluna da direita
        cw = R - CX
        kicker(s, CX, 1.72, cw, "O SETUP")
        tf = caixa(s, CX, 2.0, cw, 1.25).text_frame
        for i, (lead, resto) in enumerate([
                ("CLAUDE.md.", "Regras e papéis, lidos no começo de toda sessão."),
                ("Ritual de fim.", "Toda sessão fecha escrevendo no LOG e nas decisões pendentes."),
                ("Pesquisa vira arquivo.", "E cita os arquivos que usou — foi isso que virou o grafo.")]):
            paragrafo(tf, [(lead + " ", SEMI, 11, TEXT), (resto, LIGHT, 11, BODY)],
                      primeiro=i == 0, entre=1.1, depois=3)

        card(s, CX, 3.35, cw, 1.42)
        kicker(s, CX + 0.26, 3.55, cw - 0.5, f"PROMPT REAL  ·  SESSÃO {trace['sessao']}  ·  08/09/2026")
        prompt = trace["prompt"].strip()
        if len(prompt) > 230:
            prompt = prompt[:230].rsplit(" ", 1)[0] + " …"
        texto(s, CX + 0.26, 3.82, cw - 0.5, 0.9, [(f"“{prompt}”", SERIF, 10.5, BODY, False, 0, True)],
              entre=1.12)

        kicker(s, CX, 4.98, cw, "O QUE ACENDEU, NA ORDEM")
        col_w = cw / 2
        for i, a in enumerate(acesos):
            cx = CX + (i // 6) * col_w
            cy = 5.26 + (i % 6) * 0.235
            t = texto(s, cx, cy, col_w - 0.1, 0.22,
                      [(f"{i + 1:02d}  ", MONO, 8.5, FOOT),
                       (a["nome"], MONO, 9, FAMILIA_COR[a["familia"]])],
                      anchor=MSO_ANCHOR.MIDDLE, wrap=False)
            t.name = f"!!li_{i}"
        if final:
            familias = []
            for a in arquivos:
                rot = FAMILIA_ROT[a["familia"]]
                if rot not in familias:
                    familias.append(rot)
            texto(s, CX, 6.72, cw, 0.3,
                  [(f"{trace['comandos']} comandos · {passos} arquivos · ", MONO, 9, TEXT, True),
                   ("  →  ".join(familias) + "  →  dado", MONO, 9, ORANGE, True)],
                  anchor=MSO_ANCHOR.MIDDLE)
        elif not acesos:
            texto(s, CX, 5.26, cw, 0.3, [("clique para ver o caminho", MONO, 9, FOOT, False, 0, True)],
                  anchor=MSO_ANCHOR.MIDDLE)

        s.notes_slide.notes_text_frame.text = (
            "Sequência Morph com avanço automático: um clique no slide de repouso toca o caminho inteiro. "
            "O caminho é o real, medido no transcript da sessão (trace_sessao.py), não desenhado.")
        if k > 0:
            morph(s, 550, None if final else PASSO_MS)
        sela(s, f"s12_{k}")
        slides.append(s)
    return slides



# ================================================================== slide 13
SUB_CURTO = {"Decisões registradas": "escaladas ao humano",
             "Taxa de recorreção": "pediram 2ª rodada"}


def dados_ia():
    """Os indicadores da página 5, lidos do CSV que `graficos_p5.py` gravou."""
    quadro = pd.read_csv(DADOS / "log_sessoes.csv", parse_dates=["data"])
    return graficos_p5.indicadores(quadro)


def s13_conteudo(s):
    """Dois gráficos do relatório lado a lado e os quatro indicadores embaixo.

    Separado da moldura para poder ser reaplicado num slide já existente do
    deck (o dono edita o .pptx à mão; `monta()` não pode sobrescrevê-lo).
    """
    we = 6.2                                  # erros à esquerda, donuts à direita
    xr = L + we + 0.43
    kicker(s, L, 2.05, we, "145 ERROS DA IA  ·  QUEM PEGOU")
    imagem(s, GRAFICOS / "p5_erros.png", L, 2.4, we)
    kicker(s, xr, 2.05, R - xr, "ONDE FOI O ESFORÇO")
    imagem(s, GRAFICOS / "p5_donut.png", xr, 2.5, R - xr)

    # --- os quatro indicadores, a largura toda, na faixa de baixo
    gw = (R - L - 3 * 0.12) / 4
    for i, (rot, (valor, sub)) in enumerate(dados_ia().items()):
        x = L + i * (gw + 0.12)
        card(s, x, 5.95, gw, 0.86)
        texto(s, x + 0.22, 5.99, 1.5, 0.6, [(valor, COND, 30, ORANGE)],
              anchor=MSO_ANCHOR.MIDDLE)
        texto(s, x + 1.35, 6.07, gw - 1.5, 0.3, [(rot.upper(), MONO, 7.5, TEXT, True, 60)],
              anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        texto(s, x + 1.35, 6.35, gw - 1.5, 0.3, [(SUB_CURTO.get(rot, sub), LIGHT, 8.5, MUTED)],
              anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    s.notes_slide.notes_text_frame.text = (
        "Gráficos da página 5 do relatório (Uteis/graficos/p5_*.png, graficos_p5.py; corte "
        "de 15/08, 91 sessões dos três LOGs). Erros: 145, 88 % pegos na própria sessão. "
        "Esforço: 31/54/15 % das sessões e 25/71/4 % dos tokens.")


def s13_ia_numeros(prs, numero="13"):
    s = moldura(prs, numero, "IA em números",
                "Medimos o uso de IA como medimos a estratégia: pelo registro.")
    s13_conteudo(s)
    sela(s, "s13")
    return s


# ================================================================== slide 16
ESTUDO = DADOS / "estudo"
NEG = RGBColor(0xC2, 0x5B, 0x54)


def dados_estudo():
    """Os quatro indicadores do slide 16, lidos dos CSV do estudo — nada à mão."""
    bh = pd.read_csv(ESTUDO / "brier_horizonte.csv")
    b0 = bh[(bh.familia == "pooled") & (bh.h == 0)].iloc[0]
    ff = pd.read_csv(ESTUDO / "fomc_vs_ff.csv").set_index("h")
    reg = pd.read_csv(ESTUDO / "murphy.csv").set_index("recorte").loc["h=0"]
    obs = pd.read_csv(ESTUDO / "observacoes.csv")
    n_fomc = int(ff.loc[0, "n"])
    es = pd.read_csv(ESTUDO / "event_study.csv")
    r2_poly = es[(es.familia == "FOMC") & (es.surpresa == "poly")].r2.mean()
    mart = pd.read_csv(ESTUDO / "martingale.csv")
    vr = mart[(mart.familia == "FOMC") & (mart.filtro == "todos")].iloc[0]
    placar = [
        (True, "Calibrado", f"b = {reg.b:.2f} · p = {reg.p_wald:.2f}"),
        (True, "Skill na véspera", f"BSS {b0.bss:.0%} vs 'não sei'"),
        (True, "Erra menos que o FF", f"{ff.loc[0, 'mae_poly']:.1f} × {ff.loc[0, 'mae_ff_dm']:.1f} bps"),
        (False, "Explica o dia do FOMC?", f"R² médio {r2_poly:.2f}"),
        (False, "Δp antecipa retorno?", "t ≈ 0 (Teste_sinal)"),
        (False, "Tendência no preço?", f"VR(5) = {vr.vr5:.2f} < 1"),
    ]
    placar = [(ok, rot, num.replace(".", ",")) for ok, rot, num in placar]
    return placar, {
        "brier": (f"{b0.brier:.3f}".replace(".", ","), "BRIER NA VÉSPERA",
                  f"IC95 [{b0.brier_lo:.3f}; {b0.brier_hi:.3f}]".replace(".", ",")),
        "skill": (f"{b0.bss:.0%}", "SKILL vs 'NÃO SEI'",
                  "1 − Brier / Brier uniforme"),
        "modal": (f"{ff.loc[0, 'acerto_modal_poly']:.0%}", "DECISÕES DO FED",
                  f"{n_fomc} de {n_fomc}, na véspera"),
        "mae": (f"{ff.loc[0, 'mae_poly']:.1f} × {ff.loc[0, 'mae_ff_dm']:.1f}".replace(".", ","),
                "BPS: POLY × FUT. FF", "Diebold-Mariano p < 0,01"),
    }, reg, obs


def s16_estudo(prs, numero="16"):
    """O estudo em um slide: duas observações (frase + número + gráfico) e uma conclusão."""
    _, kpis, _, obs = dados_estudo()
    s = moldura(prs, numero, "O Polymarket como fonte de dados",
                "Calibrado e mais preciso que o mercado de juros.")

    cw = (R - L - 0.5) / 2                      # duas colunas
    observacoes = [
        ("OBSERVAÇÃO 1  ·  CALIBRAÇÃO",
         "Quando o Polymarket diz X %, acontece X % das vezes.",
         kpis["brier"][0], "Brier na véspera", kpis["brier"][2],
         "estudo_s0_calibracao.png"),
        ("OBSERVAÇÃO 2  ·  DECISÃO DO FED",
         "Na decisão do Fed, o Polymarket erra 0,9 bps; o mercado de juros, 4,8.",
         kpis["mae"][0], "bps de erro na véspera", "Polymarket × futuro de FF",
         "estudo_s2_mae.png"),
    ]
    for i, (rot, frase, numero_, sub1, sub2, arquivo) in enumerate(observacoes):
        x = L + i * (cw + 0.5)
        kicker(s, x, 1.55, cw, rot)
        texto(s, x, 1.8, cw, 0.62, [(frase, SEMI, 14, TEXT)], anchor=MSO_ANCHOR.TOP, entre=1.05)
        # linha do número (número grande + rótulos ao lado) e o gráfico na largura toda
        texto(s, x, 2.42, 1.7, 0.62, [(numero_, COND, 34, ORANGE)], anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        xn = x + (1.75 if len(numero_) > 5 else 1.25)
        texto(s, xn, 2.47, cw - 1.8, 0.26, [(sub1.upper(), MONO, 7.5, TEXT, True, 40)],
              anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        texto(s, xn, 2.73, cw - 1.8, 0.26, [(sub2, LIGHT, 9, MUTED)], anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        imagem(s, GRAFICOS / arquivo, x, 3.12, cw)

    # --- a conclusão, num card de largura inteira
    card(s, L, 6.12, R - L, 0.72)
    kicker(s, L + 0.25, 6.18, 3.0, "CONCLUSÃO", cor=MINT)
    texto(s, L + 0.25, 6.37, R - L - 0.5, 0.42,
          [("É uma probabilidade confiável — e melhor que os juros no que os dois precificam. ",
            SEMI, 13, TEXT),
           ("Por isso ela é a opinião do Black-Litterman.", SEMI, 13, ORANGE)],
          anchor=MSO_ANCHOR.MIDDLE, wrap=False)

    n_c = int(obs.groupby(["familia", "mercado", "evento"]).ngroups)
    texto(s, L, 6.93, R - L, 0.16,
          [(f"{n_c} contratos do Polymarket (Fed e CPI) · {obs.evento.nunique()} eventos · "
            f"{len(obs)} leituras em 11 horizontes · resolução no FRED · apêndice A1–A3",
            MONO, 7.5, MUTED)], wrap=False)
    s.notes_slide.notes_text_frame.text = (
        "Fonte: scripts/estudo_polymarket.py → Uteis/analises/Metricas_estudo.md; guia de "
        "leitura em Final/estudo/Guia_slide16.md. Amostra = todo o dado macro do repo "
        "(FOMC 76 faixas/18 reuniões, CPI 74 faixas/12 meses), sem exclusão. Brier 0,043 "
        "[0,022; 0,067]; calibração b = 1,04 (p = 0,15); FF = proxy DTB3−DFF corrigido do "
        "viés (cru: 8,9 bps), DM p = 0,002; 17/17 no desfecho modal. Terceira informação, "
        "se pedirem: o movimento do preço não é sinal (VR(5) = 0,78 < 1, Δp com t ≈ 0) — "
        "por isso usamos o nível, não o Δp. Detalhe nos apêndices A1–A3.")
    sela(s, "s16")
    return s


def sA1_horizonte(prs, numero="A1"):
    """Apêndice: o erro cai até o evento e o quanto o preço acrescenta ao 'não sei'."""
    s = moldura(prs, numero, "Apêndice · acurácia por horizonte",
                "Quanto mais perto do evento, menor o erro — e na véspera está no nível do Kalshi.")
    w = 9.6
    kicker(s, L + (R - L - w) / 2, 1.55, w, "BRIER POR DIAS ANTES DA RESOLUÇÃO  ·  SKILL SOBRE A PMF UNIFORME")
    imagem(s, GRAFICOS / "estudo_2_brier.png", L + (R - L - w) / 2, 1.8, w)
    frases = [
        "FOMC: 0,035 a 20 dias → 0,007 na véspera. CPI: 0,092 → 0,078 (é o objeto difícil; "
        "ninguém sabe o CPI na véspera).",
        "A linha branca (as duas famílias) para em 20 dias porque os mercados de CPI vivem "
        "~30 dias; depois só sobra FOMC.",
    ]
    for i, f in enumerate(frases):
        texto(s, L, 6.05 + i * 0.32, R - L, 0.28, [(f, LIGHT, 9.5, BODY)],
              anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    s.notes_slide.notes_text_frame.text = (
        "Tabela: Uteis/dados/estudo/brier_horizonte.csv (IC por bootstrap de evento). "
        "BSS = 1 − Brier/Brier da uniforme (1/K por faixa): 0,73 pooled na véspera; "
        "FOMC 0,96, CPI 0,42. A 60 dias o FOMC ainda tem BSS 0,67.")
    sela(s, "sA1")
    return s


def sA2_poly_vs_ff(prs, numero="A2"):
    """Apêndice: Polymarket × proxy do futuro de FF por horizonte e o lead-lag."""
    s = moldura(prs, numero, "Apêndice · Polymarket × mercado de juros",
                "No FOMC, o Polymarket erra menos e o futuro de FF não acrescenta informação.")
    kicker(s, L, 1.55, 7.2, "ERRO DO Δ TAXA E ACERTO DO DESFECHO MODAL, POR HORIZONTE")
    imagem(s, GRAFICOS / "estudo_3_fomc_ff.png", L, 1.8, 7.2)
    xr = L + 7.2 + 0.3
    kicker(s, xr, 1.55, R - xr, "QUEM SE MOVE PRIMEIRO")
    imagem(s, GRAFICOS / "estudo_4_leadlag.png", xr, 1.8, R - xr)
    enc = pd.read_csv(ESTUDO / "encompassing.csv").set_index("h")
    frases = [
        (f"Encompassing na véspera:  Δreal = a + {enc.loc[0, 'b_poly']:.2f}·E_poly "
         f"{enc.loc[0, 'b_ff']:+.2f}·E_FF   (t = {enc.loc[0, 't_poly']:.1f} e "
         f"{enc.loc[0, 't_ff']:.1f}).").replace(".", ","),
        "Proxy = DTB3 − DFF (sem o contrato ZQ, D12), corrigido do viés pela média expansiva "
        "dos erros passados (D12a).",
        "Lead-lag: nos lags sem sobreposição de janela (k = +1, −2) a correlação é ≈ 0 — a "
        "informação chega aos dois no mesmo pregão.",
    ]
    for i, f in enumerate(frases):
        texto(s, L, 5.35 + i * 0.32, R - L, 0.28, [(f, LIGHT, 9.5, BODY)],
              anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    s.notes_slide.notes_text_frame.text = (
        "Tabelas: Uteis/dados/estudo/fomc_vs_ff.csv, encompassing.csv, leadlag_*.csv. "
        "DM com correção HLN; n = 17-18 reuniões por horizonte. A 60 dias a vantagem some "
        "(p = 0,76). Granger com Newey-West: t = 1,8 (poly → FF) e 1,9 (FF → poly).")
    sela(s, "sA2")
    return s


def sA3_aplicabilidade(prs, numero="A3"):
    """Apêndice: event-study do dia do anúncio e o placar 'onde está o valor'."""
    s = moldura(prs, numero, "Apêndice · onde está o valor",
                "O nível é sinal; o movimento não é — por isso Black-Litterman, e não trading de Δp.")
    wl = 6.3
    kicker(s, L, 1.55, wl, "QUAL EXPECTATIVA EXPLICA O RETORNO DO DIA DO FOMC (R², 9 ETFs)")
    imagem(s, GRAFICOS / "estudo_5_event_study.png", L, 1.8, wl)
    xr = L + wl + 0.35
    kicker(s, xr, 1.55, R - xr, "PLACAR DO ESTUDO")
    imagem(s, GRAFICOS / "estudo_6_sintese.png", xr, 1.8, R - xr)
    frases = [
        "Na véspera quase não sobra surpresa (MAE 0,9 bps): o que move os ETFs no dia é o "
        "caminho da política, não a decisão.",
        "Antes do anúncio, Δp não antecipa o retorno (Teste_sinal.md: t ≈ 0 em h = 0, 1, 5); "
        "o preço do Polymarket reverte em 12 h (ACF −0,13, VR(5) = 0,78).",
    ]
    for i, f in enumerate(frases):
        texto(s, L, 5.35 + i * 0.32, wl + 0.4, 0.28, [(f, LIGHT, 9.5, BODY)],
              anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    s.notes_slide.notes_text_frame.text = (
        "Tabelas: Uteis/dados/estudo/event_study.csv e martingale.csv. FOMC n = 17; a "
        "surpresa-FF corrigida explica mais (R² médio 0,19 vs 0,08) porque carrega o caminho "
        "de 3 meses, não porque o Polymarket erre. CPI: R² ≈ 0 para as duas. VR(5) por "
        "família: FOMC 0,78, CPI 0,72, payrolls 0,43 — nenhuma > 1.")
    sela(s, "sA3")
    return s


# =============================================================== slides 14/15/A9
ANALISE = DADOS / "analise_backtest"


def br(v, f="{:.1f}"):
    """Número no formato do deck: vírgula decimal e sinal de menos tipográfico."""
    return f.format(v).replace(".", ",").replace("-", "–")


def dados_analise():
    """Tudo que os slides 14, 15 e A9 mostram, lido dos CSV de `analise_backtest.py`."""
    inf = pd.read_csv(ANALISE / "inferencia.csv").set_index("métrica")["valor"]
    dsr = pd.read_csv(ANALISE / "dsr.csv")
    conc = pd.read_csv(ANALISE / "concentracao.csv").set_index("métrica")["valor"]
    met = pd.read_csv(ANALISE / "metades.csv", index_col=0)
    reg = pd.read_csv(ANALISE / "regime.csv").set_index("regime")
    est = pd.read_csv(ANALISE / "estabilidade.csv").set_index("métrica")["valor"]
    imp = pd.read_csv(ANALISE / "implementacao.csv").set_index("métrica")["valor"]
    dd = pd.read_csv(ANALISE / "drawdowns.csv")
    f = lambda k: float(inf[k])
    ic_ = lambda k: inf[k].replace(".", ",").replace("-", "−")
    T = int(f("pregões (T)"))
    k = {
        "T": T,
        "excesso": br(f("excesso acumulado (pp)")), "ic_exc": ic_("excesso · IC95 bootstrap (pp)"),
        "p_exc": f"{f('P(excesso > 0)'):.0%}",
        "sharpe": br(f("Sharpe Kairós"), "{:.2f}"), "ic_sr": ic_("Sharpe Kairós · IC95 bootstrap"),
        "sharpe_spy": br(f("Sharpe SPY"), "{:.2f}"), "ic_dsr": ic_("ΔSharpe · IC95 bootstrap"),
        "p_dsr": f"{f('P(ΔSharpe > 0)'):.0%}",
        "alpha": br(f("alpha anualizado") * 100, "{:+.1f}"), "t_alpha": br(f("t do alpha (Newey-West)"), "{:.1f}"),
        "beta": br(f("beta vs SPY"), "{:.2f}"), "ir": br(f("information ratio"), "{:.2f}"),
        "te": br(f("tracking error") * 100), "psr": f"{f('PSR(SR* = 0)'):.0%}",
        "mintrl": int(round(f("MinTRL a 95 % (pregões)"))),
        "skew": br(f("assimetria"), "{:.2f}"), "kurt": br(f("curtose"), "{:.0f}"),
        "acf": br(f("ACF(1) dos retornos"), "{:.2f}"),
        "hit": f"{float(conc['hit ratio do tilt']):.0%}", "payoff": br(float(conc["payoff (ganho / |perda|)"]), "{:.2f}"),
        "hhi": br(float(conc["HHI dos dias positivos"]), "{:.3f}") + " / " + br(float(conc["HHI dos dias negativos"]), "{:.3f}"),
        "tilt": br(float(conc["soma do tilt (pp)"])), "top3": br(float(conc["3 maiores dias (pp)"])),
        "sem_top3": br(float(conc["sem os 3 maiores (pp)"])), "sem_piores3": br(float(conc["sem os 3 piores (pp)"])),
        "exc_m1": br(met.loc["1ª metade", "excesso (pp)"], "{:+.1f}"), "exc_m2": br(met.loc["2ª metade", "excesso (pp)"], "{:+.1f}"),
        "alta_bps": br(reg.loc["SPY em alta", "excesso médio (bps)"], "{:+.1f}"),
        "queda_bps": br(reg.loc["SPY em queda", "excesso médio (bps)"], "{:+.1f}"),
        "alta_hit": f"{reg.loc['SPY em alta', 'hit do excesso']:.0%}", "queda_hit": f"{reg.loc['SPY em queda', 'hit do excesso']:.0%}",
        "dd": br(float(est["máx. drawdown Kairós (%)"])), "dd_spy": br(float(est["máx. drawdown SPY (%)"])),
        "tuw": int(float(est["time under water Kairós (pregões)"])), "tuw_spy": int(float(est["time under water SPY (pregões)"])),
        "meses": est["meses K > SPY"], "roll_pos": f"{float(est['rolling IR 63d > 0 (fração)']):.0%}",
        "roll_minmax": est["rolling IR 63d mín / máx"].replace(".", ",").replace("-", "−"),
        "beta_minmax": est["rolling beta 63d mín / máx"].replace(".", ",").replace("-", "−"),
        "giro": br(float(imp["giro diário médio"]), "{:.2f}"), "desfeito": f"{float(imp['giro desfeito em 1–2 pregões']):.0%}",
        "breakeven": br(float(imp["breakeven (bps por lado)"])), "custo": br(float(imp["custo pago (pp)"])),
        "dd_tab": dd,
    }
    med = dsr[dsr["variância"].str.startswith("V medido")].set_index("N")["DSR"]
    ind = dsr[dsr["variância"].str.startswith("V de tent")].set_index("N")["DSR"]
    k["dsr_ns"] = list(med.index)
    k["dsr_med"] = f"{med.min():.2f}–{med.max():.2f}".replace(".", ",")
    k["dsr_ind"] = f"{ind.min():.2f}–{ind.max():.2f}".replace(".", ",")
    return k


def placar_backtest():
    """Os quatro números do placar e a cascata, da série diária — nada copiado à mão."""
    from graficos_p4 import metricas
    diario = pd.read_csv(DADOS / "backtest_diario.csv", parse_dates=["data"])
    diario = diario[diario.cenario == "tilt ≤ 1"].set_index("data").sort_index()
    m = metricas(diario)
    pct = lambda v: f"{v * 100:+.1f}%".replace(".", ",").replace("-", "–")
    tiles = [
        ("RETORNO LÍQUIDO", pct(m.loc["Retorno líquido", "Kairós"]), "SPY " + pct(m.loc["Retorno líquido", "SPY"])),
        ("SHARPE", br(m.loc["Sharpe (excesso zero)", "Kairós"], "{:.2f}"), "SPY " + br(m.loc["Sharpe (excesso zero)", "SPY"], "{:.2f}")),
        ("VOLATILIDADE", pct(m.loc["Vol. anualizada", "Kairós"]).lstrip("+"), "SPY " + pct(m.loc["Vol. anualizada", "SPY"]).lstrip("+")),
        ("MÁX. QUEDA", pct(m.loc["Máx. drawdown", "Kairós"]), "SPY " + pct(m.loc["Máx. drawdown", "SPY"])),
    ]
    soma = lambda c: float(diario[c].sum()) * 100
    cascata = {"mercado": soma("r_mercado"), "views": soma("r_tilt"), "custo": -soma("custo"),
               "entregue": soma("r_liquido")}
    periodo = f"{len(diario)} pregões · {diario.index[0]:%d/%m/%Y} a {diario.index[-1]:%d/%m/%Y} · líquido de custo"
    return tiles, cascata, periodo


def _tile(s, x, y, w, h, rotulo, numero, sub, cor_num=ORANGE, tam=24):
    card(s, x, y, w, h)
    texto(s, x + 0.14, y + 0.13, w - 0.28, 0.14, [(rotulo, MONO, 7, ORANGE, True, 120)], wrap=False)
    texto(s, x + 0.14, y + 0.31, w - 0.28, 0.39, [(numero, COND, tam, cor_num)], anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    texto(s, x + 0.14, y + 0.7, w - 0.28, 0.17, [(sub, BAHN, 8, MUTED)], anchor=MSO_ANCHOR.MIDDLE, wrap=False)


def s14_resultados(prs, numero="14"):
    """O slide 28 do deck da final, com UMA análise a mais: a probabilidade de o excesso ser real."""
    k = dados_analise()
    tiles, cascata, periodo = placar_backtest()
    s = moldura(prs, numero, "Backtest e resultados", f"{k['excesso']} pp sobre o SPY, ao mesmo risco")

    # --- esquerda: o placar 2 × 2 (igual ao deck), o tile novo, a cascata
    wl = 3.1
    kicker(s, L, 1.42, wl, "O PLACAR")
    tw, th = (wl - 0.14) / 2, 0.94
    for i, (rot, num, sub) in enumerate(tiles):
        _tile(s, L + (i % 2) * (tw + 0.14), 1.65 + (i // 2) * (th + 0.09), tw, th, rot, num, sub)
    # o único acréscimo: sorte ou habilidade?
    card(s, L, 3.71, wl, th)
    texto(s, L + 0.14, 3.84, wl - 0.28, 0.14, [("P(EXCESSO > 0)  ·  SORTE OU HABILIDADE?", MONO, 7, MINT, True, 120)], wrap=False)
    texto(s, L + 0.14, 4.02, 1.1, 0.39, [(k["p_exc"], COND, 24, MINT)], anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    texto(s, L + 1.15, 4.0, wl - 1.3, 0.2, [(f"IC95 {k['ic_exc']} pp · bootstrap", BAHN, 8, TEXT)], anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    texto(s, L + 1.15, 4.2, wl - 1.3, 0.2, [(f"faltam {k['mintrl'] - k['T']} pregões p/ 95 % (MinTRL)", BAHN, 8, MUTED)], anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    texto(s, L + 0.14, 4.41, wl - 0.28, 0.2, [("positivo em toda leitura; ainda não é prova", BAHN, 7.2, MUTED, False, 0, True)], wrap=False)
    texto(s, L, 4.72, wl, 0.14, [("placar completo no apêndice A8 · estatística no A9", BAHN, 7.2, MUTED)], wrap=False)

    kicker(s, L, 4.95, wl, "DE ONDE VEM O RESULTADO")
    imagem(s, GRAFICOS / "p4_composicao.png", L, 5.17, wl)
    texto(s, L, 6.78, wl, 0.34,
          [(f"Soma dos retornos diários, não capitalização — daí {br(cascata['entregue'], '{:+.1f}')} pp aqui e "
            f"{tiles[0][1]} na curva. Sem juros: taxa livre zero para os dois.",
            BAHN, 7.2, MUTED)], entre=1.0)

    # --- direita: a curva, como no deck
    xr = L + wl + 0.35
    fio(s, xr - 0.18, 1.42, xr - 0.18, 6.95)
    kicker(s, xr, 1.42, 4.0, "CURVA ACUMULADA E DRAWDOWN")
    texto(s, R - 4.4, 1.63, 4.4, 0.17, [(periodo, MONO, 7.4, MUTED)], alinha=PP_ALIGN.RIGHT, wrap=False)
    imagem(s, GRAFICOS / "p4_curva.png", xr, 1.82, R - xr)

    s.notes_slide.notes_text_frame.text = (
        "Exposição à queda. Nos dois tombos a carteira caiu junto e, em abril/2025, mais fundo que o índice: "
        "−19,6% contra −18,8%. O ganho vem da subida, não da defesa.\n"
        "Uma janela, um regime. 374 pregões de alta do S&P, sem correção para as ~200 comparações da busca "
        "tática. O teste em regime de queda é o que falta.\n"
        "Teto de risco: somar é dividir. A camada tática faz +42,6 pp sozinha e −2,94 pp dentro da carteira. "
        "Sob teto fixo, entrar não acrescenta — reparte o mesmo orçamento.\n"
        "Giramos os quatro botões da estratégia — risco, confiança, gatilho e correção das probabilidades — e a "
        "vantagem sobre o SPY não vira negativa em nenhuma posição: +0,5 a +15,8 pp. Não é um parâmetro bem "
        "escolhido.\n"
        f"O tile novo (sorte ou habilidade?): bootstrap estacionário (Politis-Romano, blocos de 10 pregões, "
        f"B = 2 000) sobre a série diária. IC95 do excesso {k['ic_exc']} pp, P(> 0) = {k['p_exc']}; Sharpe "
        f"{k['sharpe']} com IC95 {k['ic_sr']}; alpha t = {k['t_alpha']} (Newey-West) contra a régua de 3 de "
        f"Harvey-Liu; PSR = {k['psr']}; track record mínimo a 95 % = {k['mintrl']} pregões, faltam "
        f"{k['mintrl'] - k['T']}. Leitura: o número é positivo em toda leitura, mas 18 meses não separam "
        "habilidade de sorte — o que sustenta a estratégia é o mecanismo medido fora do resultado (slides 8, 9 "
        "e 16). Detalhe no apêndice A9 e em Uteis/analises/Analise_backtest.md.")
    sela(s, "s14")
    return s


def s15_analise_critica(prs, numero="15"):
    """Três números, uma leitura qualitativa: ainda não provada · defensiva · rigorosa onde depende de nós."""
    k = dados_analise()
    s = moldura(prs, numero, "Análise crítica",
                "Três números, uma leitura: positiva, defensiva — e ainda não provada")
    cw = (R - L - 0.6) / 3
    colunas = [
        ("A PROVA EXIGE TEMPO", str(k["mintrl"]), f"PREGÕES DE TRACK RECORD MÍNIMO A 95 %  ·  TEMOS {k['T']}",
         "O excesso é positivo em toda leitura, mas ainda não é prova de habilidade. "
         f"Faltam {k['mintrl'] - k['T']} pregões — o próximo semestre é o teste fora da amostra, "
         "pré-registrado antes de ser visto.",
         f"PSR {k['psr']} · IC95 do excesso {k['ic_exc']} pp · alpha t = {k['t_alpha']} contra a régua de 3 (Harvey–Liu)"),
        ("ONDE O GANHO MORA", f"{k['queda_bps']} bps", f"POR DIA DE QUEDA DO SPY  ·  {k['alta_bps']} NOS DE ALTA",
         "Perde menos quando o mercado cai, ganha menos quando sobe: é um tilt defensivo sobre beta "
         f"{k['beta']}, não proteção de cauda — no tombo de abril caiu mais fundo que o índice, e na "
         "2ª metade, só de alta, ficou atrás.",
         f"hit {k['queda_hit']} nos dias de queda, {k['alta_hit']} nos de alta · metades {k['exc_m1']} / {k['exc_m2']} pp · "
         f"máx. queda {k['dd']} % × SPY {k['dd_spy']} %"),
        ("PELO PROTOCOLO DA LITERATURA", "3 de 7", "PONTOS DE ARNOTT, HARVEY & MARKOWITZ FECHADOS",
         "Rigor onde depende de nós: hipótese antes do teste, dados sem look-ahead, admissão por mecanismo. "
         "Parcial onde o desenho limita: testes múltiplos, Ω diagonal. Aberto onde só o tempo resolve: "
         "validação fora da amostra.",
         "fechados: motivação · dados · cultura  —  parciais: testes múltiplos · dinâmica · complexidade  —  aberto: fora da amostra"),
    ]
    y0, h = 1.62, 3.9
    for i, (rot, num, sub, conclusao, evid) in enumerate(colunas):
        x = L + i * (cw + 0.3)
        card(s, x, y0, cw, h)
        kicker(s, x + 0.25, y0 + 0.22, cw - 0.5, rot)
        texto(s, x + 0.25, y0 + 0.5, cw - 0.5, 0.8, [(num, COND, 44, ORANGE)], anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        texto(s, x + 0.25, y0 + 1.34, cw - 0.5, 0.36, [(sub, MONO, 7.5, TEXT, True, 30)], entre=1.05)
        texto(s, x + 0.25, y0 + 1.82, cw - 0.5, 1.7, [(conclusao, SEMI, 11, TEXT)], entre=1.08)
        texto(s, x + 0.25, y0 + 3.15, cw - 0.5, 0.55, [(evid, LIGHT, 8.5, MUTED)], entre=1.05)

    kicker(s, L, 5.9, 3.0, "FARÍAMOS DIFERENTE", cor=MINT)
    texto(s, L, 6.12, R - L, 0.24,
          [("Pré-registrar o próximo semestre como teste fora da amostra · corrigir a busca tática pelo nº de "
            "tentativas (DSR) · um Ω que enxergue a correlação entre views.", LIGHT, 10, BODY)], wrap=False)
    texto(s, L, 6.93, R - L, 0.16,
          [("Arnott, Harvey & Markowitz (2019) · Bailey & López de Prado (2012, 2014) · Harvey & Liu (2015)  ·  "
            "Final/analise/Pesquisa_avaliacao_estrategia.md  ·  apêndice A9", MONO, 7.5, MUTED)], wrap=False)
    s.notes_slide.notes_text_frame.text = (
        "Fala (~60 s): três números para ler a estratégia. (1) 457: o track record mínimo para dizer com 95 % que "
        f"o Sharpe é positivo — temos {k['T']}. O excesso é positivo em toda leitura (P = {k['p_exc']}), mas não é "
        "prova; o próximo semestre é o teste fora da amostra, e vai ser pré-registrado. (2) +6,2 bps por dia de "
        "queda do SPY, −3,9 nos de alta: o ganho vem de perder menos quando o mercado cai — tilt defensivo sobre "
        "beta 0,95. Não é proteção de cauda: no tombo de abril/2025 caímos mais fundo (−19,6 % × −18,8 %) e na 2ª "
        "metade, só de alta, ficamos atrás (−0,6 pp). (3) 3 de 7 no protocolo de Arnott, Harvey e Markowitz: "
        "fechados motivação econômica, dados/amostra e cultura de pesquisa; parciais testes múltiplos (~200 "
        "comparações da tática sem correção; DSR reportado no A9), dinâmica (parâmetros fixos, mas excesso só na "
        "1ª metade) e complexidade (Ω diagonal não vê correlação entre views); aberto validação fora da amostra.\n"
        "Se perguntarem 'por que mostrar os fracos?': porque o protocolo pede e porque um slide só de fortes em "
        "18 meses de dado é o que a literatura chama de storytelling. Scorecard completo em "
        "Final/analise/Pesquisa_avaliacao_estrategia.md §2.")
    sela(s, "s15")
    return s


def sA9_estatistica(prs, numero="A9"):
    """Apêndice: a tear-sheet (López de Prado, cap. 14) e as duas curvas que a sustentam."""
    k = dados_analise()
    s = moldura(prs, numero, "Apêndice · estatística do backtest",
                "Tear-sheet completa: inferência, concentração, estabilidade e implementação")
    wl = 5.7
    kicker(s, L, 1.55, wl, f"TILT ≤ 1  ·  {k['T']} PREGÕES  ·  IC95 BOOTSTRAP")
    dd = k["dd_tab"]
    linhas = [
        ("INFERÊNCIA", None),
        ("Sharpe Kairós [IC95]", f"{k['sharpe']}  {k['ic_sr']}"),
        ("Sharpe SPY  ·  ΔSharpe [IC95]", f"{k['sharpe_spy']}  ·  {k['ic_dsr']}  P = {k['p_dsr']}"),
        ("excesso acumulado [IC95]", f"{k['excesso']} pp  {k['ic_exc']}  P = {k['p_exc']}"),
        ("alpha anual (t Newey-West)  ·  beta", f"{k['alpha']} %  (t {k['t_alpha']})  ·  {k['beta']}"),
        ("information ratio  ·  tracking error", f"{k['ir']}  ·  {k['te']} %"),
        ("PSR(SR > 0)  ·  MinTRL 95 %", f"{k['psr']}  ·  {k['mintrl']} pregões"),
        ("DSR, N = " + " / ".join(str(n) for n in k["dsr_ns"]), f"V medido {k['dsr_med']}  ·  V indep. {k['dsr_ind']}"),
        ("assimetria  ·  curtose  ·  ACF(1)", f"{k['skew']}  ·  {k['kurt']}  ·  {k['acf']}"),
        ("CONCENTRAÇÃO (PERNA DAS VIEWS)", None),
        ("hit ratio  ·  payoff  ·  HHI +/−", f"{k['hit']}  ·  {k['payoff']}  ·  {k['hhi']}"),
        ("tilt  ·  sem os 3 maiores  ·  sem os 3 piores", f"{k['tilt']}  ·  {k['sem_top3']}  ·  {k['sem_piores3']} pp"),
        ("ESTABILIDADE", None),
        ("excesso 1ª / 2ª metade", f"{k['exc_m1']} / {k['exc_m2']} pp"),
        ("excesso por dia: SPY em alta / em queda", f"{k['alta_bps']} / {k['queda_bps']} bps  (hit {k['alta_hit']} / {k['queda_hit']})"),
        ("máx. drawdown  ·  time under water", f"{k['dd']} % ({k['tuw']} pregões)  ·  SPY {k['dd_spy']} % ({k['tuw_spy']})"),
        ("maior drawdown", f"{dd.loc[0, 'início']} → {dd.loc[0, 'fundo']} → {dd.loc[0, 'recuperação']}"),
        ("meses > SPY  ·  IR 63d > 0  ·  beta 63d", f"{k['meses']}  ·  {k['roll_pos']}  ·  {k['beta_minmax']}"),
        ("IMPLEMENTAÇÃO", None),
        ("giro/dia  ·  desfeito em 1–2 pregões", f"{k['giro']}  ·  {k['desfeito']}"),
        ("custo pago  ·  breakeven", f"{k['custo']} pp  ·  {k['breakeven']} bps por lado"),
    ]
    y = 1.82
    for rot, val in linhas:
        if val is None:
            fio(s, L, y + 0.05, L + wl, y + 0.05, alfa=60000)
            kicker(s, L, y + 0.09, wl, rot, cor=MUTED)
            y += 0.3
            continue
        texto(s, L, y, 2.75, 0.2, [(rot, LIGHT, 8.5, BODY)], anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        texto(s, L + 2.8, y, wl - 2.8, 0.2, [(val, MONO, 8, TEXT)], anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        y += 0.215

    xr, wr = L + wl + 0.4, R - L - wl - 0.4
    kicker(s, xr, 1.55, wr, "INFORMATION RATIO EM JANELA MÓVEL DE 63 PREGÕES")
    imagem(s, GRAFICOS / "ab_rolling_ir.png", xr, 1.78, wr)
    kicker(s, xr, 4.2, wr, "DEFLATED SHARPE RATIO PELO Nº DE TENTATIVAS")
    imagem(s, GRAFICOS / "ab_dsr.png", xr + 0.3, 4.43, wr - 1.7)
    texto(s, xr, 6.58, wr, 0.36,
          [("Laranja: variância dos Sharpe das configurações gravadas (correlacionadas — subestima). "
            "Azul: tentativas independentes (cota superior). A verdade fica entre as duas.", LIGHT, 8.5, BODY)], entre=1.05)
    texto(s, L, 6.93, R - L, 0.16,
          [("Tabelas em Uteis/dados/analise_backtest/  ·  scripts/analise_backtest.py --demo  ·  "
            "retornos mensais e os 3 maiores drawdowns em Uteis/analises/Analise_backtest.md", MONO, 7.5, MUTED)], wrap=False)
    s.notes_slide.notes_text_frame.text = (
        "Definições: PSR = P(Sharpe verdadeiro > 0) dado T, assimetria e curtose (Bailey-LdP 2012). "
        "MinTRL = T necessário para PSR = 95 %. DSR = PSR contra o Sharpe esperado do máximo de N tentativas "
        "de ruído (Bailey-LdP 2014); N vem do registro do projeto (configurações gravadas, hipóteses do LOG, "
        "células da D27). HHI = concentração de Herfindahl dos dias positivos/negativos (0 = uniforme). "
        "TuW = maior sequência de pregões abaixo do pico. Bootstrap: Politis-Romano, bloco médio 10, "
        "B = 2 000, semente 20260918. Alpha com Newey-West 5 lags.")
    sela(s, "sA9")
    return s


# ============================================================== transição
def sT_resposta(prs, numero=""):
    """Slide de transição entre a análise crítica e o estudo: a hipótese volta ao palco."""
    s = moldura(prs, numero, "Resposta da hipótese", "")
    texto(s, L, 2.25, R - L, 0.9, [("Resposta da hipótese", COND, 54, TEXT)], anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    kicker(s, L, 3.45, 6.0, "A HIPÓTESE, DO SLIDE 2")
    texto(s, L, 3.72, R - L - 1.5, 1.1,
          [("“Mercados de previsão antecipam o mercado financeiro. ", SEMI, 22, BODY),
           ("Esse é o nosso sinal.”", SEMI, 22, ORANGE)], entre=1.1)
    obs = pd.read_csv(ESTUDO / "observacoes.csv")
    n_c = int(obs.groupby(["familia", "mercado", "evento"]).ngroups)
    texto(s, L, 5.15, R - L, 0.3,
          [(f"O que {n_c} contratos do Polymarket, {obs.evento.nunique()} eventos e {len(obs)} leituras dizem.",
            LIGHT, 14, MUTED)], wrap=False)
    s.notes_slide.notes_text_frame.text = (
        "Transição (~10 s): 'Voltamos à hipótese do começo. Mercados de previsão antecipam o mercado "
        "financeiro — esse era o sinal. A pergunta agora é se ele se confirmou no dado.' Avança para o "
        "slide do estudo (calibração e decisão do Fed).")
    sela(s, "sT")
    return s


def inserir_apos(saida: Path, construtor, prefixo: str) -> Path:
    """Acrescenta UM slide ao deck salvo, logo depois do último cujo cabeçalho começa por `prefixo`."""
    prs = Presentation(str(saida))
    n = len(prs.slides)
    pos = _indice(prs, prefixo)
    construtor(prs)
    _mover(prs, n, pos + 1)
    prs.save(str(saida))
    return saida


# ------------------------------------------------------------------- render
def render(pptx: Path, pasta: Path):
    """Exporta cada slide em PNG pelo próprio PowerPoint (COM) — conferência."""
    import win32com.client
    pasta.mkdir(parents=True, exist_ok=True)
    app = win32com.client.Dispatch("PowerPoint.Application")
    pres = app.Presentations.Open(str(pptx), WithWindow=False)
    try:
        for i, sl in enumerate(pres.Slides, 1):
            sl.Export(str(pasta / f"slide{i:02d}.png"), "PNG", 1920, 1080)
    finally:
        pres.Close()
    return sorted(pasta.glob("slide*.png"))


def monta(saida: Path = SAIDA) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    s3_polymarket(prs)
    s4_comparacao(prs)
    s12_second_brain(prs)
    s13_ia_numeros(prs)
    s14_resultados(prs)
    s15_analise_critica(prs)
    sT_resposta(prs)
    s16_estudo(prs)
    sA1_horizonte(prs)
    sA2_poly_vs_ff(prs)
    sA3_aplicabilidade(prs)
    sA9_estatistica(prs)
    saida.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(saida))
    return saida


def _indice(prs, prefixo):
    """Posição (0-based) do último slide cujo cabeçalho de seção começa por `prefixo`."""
    achado = None
    for i, sl in enumerate(prs.slides):
        if any(sh.has_text_frame and sh.text_frame.text.startswith(prefixo) for sh in sl.shapes):
            achado = i
    return achado


def _mover(prs, de, para):
    lst = prs.slides._sldIdLst
    el = list(lst)[de]
    lst.remove(el)
    lst.insert(para, el)


def anexar(saida: Path = SAIDA) -> Path:
    """Acrescenta 14, 15 e A9 ao deck JÁ SALVO, sem regenerar o resto.

    O `Slides_novos.pptx` tem edições feitas à mão no PowerPoint (sessão 47) e
    slides colados da semi; `monta()` apagaria tudo isso. Aqui os três slides
    novos entram no deck existente: 14 e 15 logo depois do "IA em números",
    A9 depois do último apêndice.
    """
    prs = Presentation(str(saida))
    n = len(prs.slides)
    # posições medidas ANTES de acrescentar (o A9 também se chama "Apêndice")
    pos_ia = _indice(prs, "KAIROS   IA em números")
    pos_ap = max(_indice(prs, p) or 0 for p in ("KAIROS   Apêndice", "KAIRÓS   Apêndice", "KAIRÓS  Apêndice"))
    s14_resultados(prs)
    s15_analise_critica(prs)
    sA9_estatistica(prs)
    _mover(prs, n, pos_ia + 1)        # 14
    _mover(prs, n + 1, pos_ia + 2)    # 15
    _mover(prs, n + 2, pos_ap + 3)    # A9: +2 pelos dois inseridos antes, +1 para ficar depois
    prs.save(str(saida))
    return saida


def demo():
    # o mapa do eixo do slide 4 é linear e cobre o alcance declarado
    assert abs(x_de(EIXO_MIN) - X0) < 1e-9 and abs(x_de(EIXO_MAX) - X1) < 1e-9
    assert fmt_pct(-0.0153) == "− 1,53%" and fmt_pct(0.0003) == "+ 0,03%"
    # o trace tem de caber no grafo: é a asserção que o slide 12 faz
    _, xy, n, _ = grafo_para_slide()
    trace = json.loads((DADOS / "trace_sessao.json").read_text(encoding="utf-8"))
    assert all(a["nome"] in xy for a in trace["arquivos"]) and n > 100
    # os indicadores do slide 16 saem do CSV do estudo e têm a forma esperada
    placar, kpis, reg, obs = dados_estudo()
    assert set(kpis) == {"brier", "skill", "modal", "mae"} and 0.9 < reg.b < 1.2
    assert len(placar) == 6 and [ok for ok, *_ in placar] == [True] * 3 + [False] * 3
    assert obs.h.nunique() == 11 and obs.familia.nunique() == 2
    # os slides 14/15/A9 leem o CSV da análise; o placar do protocolo soma 7
    k = dados_analise()
    assert k["T"] == 374 and k["ic_exc"].count(";") == 1 and k["mintrl"] > k["T"]
    tiles, cascata, _ = placar_backtest()
    assert len(tiles) == 4 and abs(cascata["mercado"] + cascata["views"] + cascata["custo"] - cascata["entregue"]) < 1e-6
    assert br(-1.25) == "–1,2"
    print("demo ok")


if __name__ == "__main__":
    if "--demo" in sys.argv:
        demo()
        sys.exit()
    if "--transicao" in sys.argv:
        saida = inserir_apos(SAIDA, sT_resposta, "KAIROS   Análise crítica")
    else:
        saida = anexar() if "--anexar" in sys.argv else monta()
    print(f"gravado {saida.relative_to(RAIZ)}")
    if "--render" in sys.argv:
        pasta = Path(sys.argv[sys.argv.index("--render") + 1]) if len(sys.argv) > sys.argv.index("--render") + 1 else saida.parent / "render"
        for p in render(saida, pasta):
            print("  ", p)

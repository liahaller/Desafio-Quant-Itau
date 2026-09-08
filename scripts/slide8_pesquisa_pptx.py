"""Slide 8 da semifinal — a peneira: so entra no modelo o que o dado provou.

Gera `Semis/Slide_8.pptx`, deck separado do `Slides_3a6.pptx` (o slide 8 nao
tem diagrama e nao participa do Morph — juntar os dois arquivos so obrigaria o
auto-teste do outro a abrir excecao para ele).

O slide e uma FIGURA, nao uma lista: a fala carrega o detalhe. Cada ponto da
fileira de cima e uma coisa que foi medida (31 no total, em quatro grupos);
embaixo deles corre a peneira — uma barra continua, furada exatamente sob quem
passou. Os 7 que passam caem na fileira de baixo e uma chave os despeja no no
do modelo. O que nao passou fica retido em cinza, na fileira de cima: a perda
e visivel sem precisar de numero escrito.

A camada visual e importada do `slides_3a6_pptx.py` (moldura, paleta, painel,
chave, flecha), para o slide entrar no meio do deck sem trocar de identidade.

FONTE DOS NUMEROS (nenhum digitado de cabeca):
  · 13 candidatas taticas medidas / 1 entrou  → `relatorio_p4.py` (p. 4 do
    relatorio final, faixa "Autoavaliacao · o placar"), D28 do `Felipe`;
  · 4 ingredientes do Omega / 2 aprovados     → mesma faixa;
  · grade de 6 janelas (1, 2, 3, 5, 10, 20)   → `premissa_g1.py::HORIZONTES`;
    o bloco k = 3, 5, 10 que sobrou           → D28.3;
  · 8 configuracoes do teto ({1,2,3,5} × dois escopos) e o nivel 1 no tilt
    escolhido PELA REGRA, antes do resultado  → D10a / D12 do `Felipe`;
  · variance ratio de 0,97 a 1,18 e as teses derrubadas → D28.0 e p. 4;
  · 9 ETFs                                    → `data/README.md`;
  · 9 mercados do Polymarket                  → D5 / `poly_preprocessing.py`;
  · 374 pregoes, 2025-02-10 a 2026-08-06      → `Uteis/analises/Backtest_v1.md`.
"""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

from relatorio_p4 import MUTED, MUTED2, ORANGE, TEXT, caixa, escreve, regua, rotulo
from slides_3a6_pptx import (BODY, EDGE, L, MINT, ON_A, ON_B, PANEL_A, PANEL_B,
                             PE, R, brilho, chave, efeito, flecha, mistura,
                             moldura, moldura_da_forma, painel, pino, sela)

RAIZ = Path(__file__).resolve().parent.parent
SAIDA = RAIZ / "Semis" / "Slide_8.pptx"

# ------------------------------------------------------------------ geometria
D, VAO, VAO_GRUPO = 0.19, 0.07, 0.46   # ponto, vao entre pontos, vao de grupo
Y_ROTULO = 1.80                        # o nome de cada grupo
Y_MEDIDAS = 2.12                       # a fileira do que foi medido
Y_PENEIRA = 2.70                       # a barra da peneira
H_PENEIRA = 0.17
Y_PASSOU = 3.18                        # a fileira do que passou
Y_CHAVE, H_CHAVE = 3.66, 0.46
Y_NO, H_NO, W_NO = 4.52, 1.00, 5.00    # o no do modelo
CAP_X = 10.50                          # a coluna de legendas, a direita
BARRADO_X = 8.35                       # o que a peneira reteve, embaixo

# --------------------------------------------------------------------- dados
# (rotulo do grupo, medidas, passaram) — a unica fonte dos pontos do slide.
#
# TODO(DECISAO-30): a 1a linha conta "13 medidas → 1 passou" (o placar da p. 4
# do relatorio) em vez do "5 candidatos, nenhum entrou" do `Semis.txt`, que
# contradiz o slide 6 e o backtest — a camada v2 ESTA ligada. A reuniao ainda
# nao escolheu como alinhar 8 e 9 ao 6; se escolher outra contagem, muda aqui.
GRUPOS = [
    ("13 ideias de sinal", 13, 1),
    ("4 critérios do Ω", 4, 2),
    ("6 janelas", 6, 3),
    ("8 tetos de risco", 8, 1),
]

# o que a peneira reteve — tres exemplos, com o teste que os derrubou
BARRADOS = [
    ("Momentum e velocidade de ajuste.",
     " Variance ratio de 0,97 a 1,18: passeio aleatório."),
    ("Proximidade do evento, no Ω.",
     " Sinal invertido nas duas views em que foi testada."),
    ("Prêmio de anúncio no Fed, event-driven.",
     " Medidas, e a tese caiu no dado — não no retorno."),
]

INSUMOS = ("9 ETFs americanos   ·   9 mercados do Polymarket   ·   "
           "taxas do FRED   ·   374 pregões (10/02/2025 a 06/08/2026)")


def posicoes():
    """`(x de cada ponto, x dos que passaram, fim da fileira)`.

    Os que passam ficam no MEIO do grupo, nao na ponta: a peneira precisa de
    barra dos dois lados de cada furo para se ler como barra furada.
    """
    todos, passaram, x = [], [], L
    for _, total, k in GRUPOS:
        inicio = (total - k) // 2
        for i in range(total):
            todos.append(x)
            if inicio <= i < inicio + k:
                passaram.append(x)
            x += D + VAO
        x += VAO_GRUPO - VAO
    return todos, passaram, x - VAO_GRUPO


def ponto(slide, x, y, aceso):
    """Um ponto da fileira: laranja e aceso se passou, cinza se ficou retido."""
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                                Inches(D), Inches(D))
    sh.fill.solid()
    sh.fill.fore_color.rgb = ORANGE if aceso else mistura(PANEL_A, EDGE, 0.55)
    sh.line.fill.background()
    if aceso:
        efeito(sh, brilho(str(ORANGE), rad=90000, alfa=38000))
    else:
        sh.shadow.inherit = False
    return sh


def peneira(slide, x0, x1, furos):
    """A barra, desenhada como os PEDACOS entre os furos.

    Um furo e a ausencia de barra — desenhar a barra inteira e depois tapar os
    furos com retangulos da cor do fundo mentiria em qualquer outro fundo.
    """
    borda, folga = mistura(EDGE, ORANGE, 0.45), 0.045
    inicio = x0
    for f in furos:
        if f - folga - inicio > 0.02:
            painel(slide, inicio, Y_PENEIRA, f - folga - inicio, H_PENEIRA,
                   topo=ON_A, base=ON_B, borda=borda, raio=0.10)
        inicio = f + D + folga
    if x1 - inicio > 0.02:
        painel(slide, inicio, Y_PENEIRA, x1 - inicio, H_PENEIRA, topo=ON_A,
               base=ON_B, borda=borda, raio=0.10)


def queda(slide, x):
    """O fio que liga o ponto de cima ao de baixo, atravessando o furo."""
    ln = slide.shapes.add_connector(1, Inches(x + D / 2), Inches(Y_MEDIDAS + D),
                                    Inches(x + D / 2), Inches(Y_PASSOU))
    ln.line.color.rgb = mistura(ORANGE, PANEL_B, 0.55)
    ln.line.width = Pt(0.75)
    return ln


def legenda(slide, y, titulo, texto=None, tamanho=8.5):
    """Rotulo curto na coluna da direita, ancorado numa altura do desenho."""
    rotulo(slide, CAP_X, y, titulo, w=R - CAP_X, tamanho=tamanho)
    if texto:
        tf = caixa(slide, CAP_X, y + 0.20, R - CAP_X, 0.34)
        escreve(tf, [(texto, MUTED, False)], 8.5, espaco=1.20, primeiro=True)


def no_do_modelo(slide, x):
    """O no de chegada: o que sobrou da peneira e o que a carteira usa."""
    sh = painel(slide, x, Y_NO, W_NO, H_NO, topo=ON_A, base=ON_B, borda=MINT,
                luz=brilho(str(MINT), rad=110000, alfa=26000))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left, tf.margin_right = Inches(0.52), Inches(0.20)
    tf.margin_top = tf.margin_bottom = 0
    escreve(tf, [("ENTRA NO MODELO", MINT, True)], 13.0, spc=0.9, espaco=1.1,
            primeiro=True)
    escreve(tf, [("as 4 views e a camada tática que a carteira roda", BODY,
                  False)], 9.5, espaco=1.1, antes=2.0)
    pino(slide, x + 0.26, Y_NO + H_NO * 0.24, H_NO * 0.52, MINT, esp=0.075)
    return sh


def s8_pesquisa(prs):
    """A peneira: 31 coisas medidas, 7 de pé, e o critério escrito antes."""
    s = moldura(prs, "08", "O que a pesquisa mediu", banda=False)
    rodape(s)

    tf = caixa(s, L, 0.96, 11.4, 0.34)
    escreve(tf, [("Para entrar no modelo, a ideia tinha que se provar no dado",
                  TEXT, True), (" — não no retorno.", BODY, False)], 19.0,
            espaco=1.0, primeiro=True)
    tf = caixa(s, L, 1.36, 11.4, 0.22)
    escreve(tf, [(INSUMOS, MUTED, False)], 10.0, primeiro=True)
    regua(s, L, 1.60, R)

    todos, passaram, fim = posicoes()

    x = L
    for nome, total, _ in GRUPOS:
        passo = total * (D + VAO) - VAO + VAO_GRUPO
        # a caixa do rotulo termina onde comeca a do grupo seguinte: mais larga
        # que isso e o rotulo do vizinho entra por baixo dela
        rotulo(s, x, Y_ROTULO, nome, w=passo - 0.06, tamanho=8.5)
        x += passo

    for f in passaram:                       # os fios vao ANTES: passam atras
        queda(s, f)
    peneira(s, L, fim, passaram)
    for px in todos:
        ponto(s, px, Y_MEDIDAS, px in passaram)
    for px in passaram:
        ponto(s, px, Y_PASSOU, True)

    legenda(s, Y_MEDIDAS - 0.06, "31 medições",
            "cada ponto é uma medição")
    legenda(s, Y_PENEIRA - 0.02, "prova estatística",
            "o critério é escrito antes de medir")
    legenda(s, Y_PASSOU - 0.08, "7 de pé", tamanho=13.0)

    esq, dir_ = passaram[0], passaram[-1] + D
    chave(s, esq, Y_CHAVE, dir_ - esq, H_CHAVE, cor=mistura(MUTED, ORANGE, 0.3))
    meio = (esq + dir_) / 2
    flecha(s, meio, Y_CHAVE + H_CHAVE, meio, Y_NO - 0.02, cor=MUTED)
    no_do_modelo(s, meio - W_NO / 2)

    # ---- o que ficou retido, ao lado do nó
    rotulo(s, BARRADO_X, Y_NO - 0.34, "O que a peneira barrou", w=R - BARRADO_X,
           tamanho=8.5)
    tf = caixa(s, BARRADO_X, Y_NO, R - BARRADO_X, 1.30)
    for i, (lead, resto) in enumerate(BARRADOS):
        escreve(tf, [(lead, TEXT, True), (resto, BODY, False)], 10.5,
                espaco=1.22, antes=0 if i == 0 else 7.0, primeiro=i == 0)

    tf = caixa(s, L, 6.32, 11.4, 0.24)
    escreve(tf, [("Reprovar era o resultado esperado.", ORANGE, True),
                 ("  Nenhum parâmetro foi escolhido olhando o retorno: janela, "
                  "horizonte e nível do teto são declarados antes da medição, e "
                  "as ~200 comparações da busca tática estão no relatório, sem "
                  "correção e sem esconder.", MUTED2, False)], 10.5,
            espaco=1.24, primeiro=True)

    s.notes_slide.notes_text_frame.text = (
        "cada ponto é uma coisa que a gente mediu: 13 ideias de sinal, os 4 "
        "critérios de confiança, 6 janelas e 8 configurações de teto de risco. "
        "A peneira é o teste estatístico, e ela só tem furo onde o dado provou "
        "a tese — 7 de 31. Momentum e velocidade de ajuste, por exemplo, "
        "morreram num variance ratio de 0,97 a 1,18: passeio aleatório. "
        "A régua é essa: a ideia entra por acertar a probabilidade, nunca por "
        "ter rendido mais.")
    sela(s, 8)
    return s


def rodape(s):
    """Reaponta o crédito do rodapé, herdado da `moldura`, para este script."""
    alvo = next(sh for sh in s.shapes if sh.has_text_frame
                and sh.text_frame.text.startswith("Diagrama e slides"))
    runs = alvo.text_frame.paragraphs[0].runs
    runs[0].text = "Cada ponto sai de artefato medido; slide gerado por "
    runs[1].text = "scripts/slide8_pesquisa_pptx.py"
    runs[2].text = "."


# ----------------------------------------------------------------- montagem
def monta(saida: Path = SAIDA) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s8_pesquisa(prs)
    saida.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(saida))
    return saida


# ---------------------------------------------------------------- auto-teste
def demo() -> None:
    """Checa moldura, `effectLst` duplicado, colisão de texto e a contagem de
    pontos. Roda com `python scripts/slide8_pesquisa_pptx.py --demo`.
    """
    from pptx import Presentation as _P

    from relatorio_p5 import altura_do_texto
    from slides_3a6_pptx import A

    todos, passaram, fim = posicoes()
    assert fim <= CAP_X - 0.15, f"a fileira invade a coluna de legendas ({fim:.2f})"
    assert len(passaram) == sum(k for _, _, k in GRUPOS)

    caminho = monta()
    prs = _P(str(caminho))
    s = prs.slides[0]
    larg, alt = prs.slide_width, prs.slide_height

    for sh in s.shapes:
        x0, y0, x1, y1 = moldura_da_forma(sh)
        assert x0 >= -1 and y0 >= -1, f"caixa fora da moldura ({sh.name})"
        assert x1 <= larg and y1 <= alt, f"caixa estoura a moldura ({sh.name})"
        sp = getattr(sh._element, "spPr", None)
        if sp is not None:
            assert len(sp.findall(A + "effectLst")) <= 1, \
                f"effectLst duplicado ({sh.name})"

    # O defeito que o .pptx NAO denuncia: um texto que ganha uma linha invade a
    # caixa de baixo e o arquivo abre sem reclamar. Mede a altura real no Segoe
    # UI e cruza todos contra todos — a altura declarada da caixa nao serve.
    # A moldura (cabecalho e rodape) fica de fora: e conferida no deck 3-7.
    caixas = []
    for sh in s.shapes:
        if not (sh.has_text_frame and sh.text_frame.text.strip()):
            continue
        y = Emu(sh.top).inches
        if y < 0.90 or y > PE or sh.shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
            continue                      # o no tem o texto DENTRO do painel
        larg_cx = Emu(sh.width).inches
        caixas.append((sh.text_frame.text[:30], Emu(sh.left).inches, y, larg_cx,
                       altura_do_texto(sh.text_frame, larg_cx)))
    for i, (na, xa, ya, wa, ha) in enumerate(caixas):
        assert ya + ha <= PE, f"texto passa do rodapé: {na!r}"
        for nb, xb, yb, wb, hb in caixas[i + 1:]:
            comum_x = min(xa + wa, xb + wb) - max(xa, xb)
            comum_y = min(ya + ha, yb + hb) - max(ya, yb)
            assert comum_x <= 0.05 or comum_y <= 0.02, \
                f"texto sobre texto: {na!r} x {nb!r}"

    pontos = sum(1 for sh in s.shapes if sh.width == Inches(D)
                 and sh.height == Inches(D))
    assert pontos == len(todos) + len(passaram), \
        f"{pontos} pontos desenhados, {len(todos)} + {len(passaram)} esperados"
    print(f"ok — {caminho.name}: {len(todos)} medidas · {len(passaram)} passam "
          f"· {len(caixas)} blocos de texto sem colisão")


if __name__ == "__main__":
    import sys

    if "--demo" in sys.argv:
        demo()
    else:
        print(monta())

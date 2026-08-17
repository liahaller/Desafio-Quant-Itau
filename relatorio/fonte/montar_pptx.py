"""Reconstrói a página 3 do relatório como um slide de PowerPoint editável.

O slide sai em 960 x 540 pt (13,333 x 7,5 in = 16:9), a mesma caixa do PDF, e as
coordenadas são as medidas no próprio `KAIROS.pdf` — texto, filetes, barras e
cartões ficam onde já estavam. Tudo é forma nativa do PowerPoint, exceto o
diagrama de arquitetura, que entra como PNG de alta resolução gerado do
`arquitetura.svg` pelo mesmo Chrome que imprime o relatório.
"""
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Pt

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

BASE = Path(__file__).resolve().parent
CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
SAIDA = BASE.parent / "KAIROS_pagina3.pptx"

# ---------------------------------------------------------------- paleta
FUNDO = "0B1622"
FUNDO2 = "101F2E"
LINHA = "24384B"
TINTA = "E8EDF2"
TINTA2 = "9BAEC0"
TINTA3 = "6B8299"
AMBAR = "F5A623"
AMBAR_D = "C9820C"
AZUL = "3D8FC4"
VERMELHO = "C2544D"

# Fundos translúcidos do CSS achatados sobre o fundo da página: o PowerPoint
# guarda alfa por forma, mas a cor composta é idêntica e sobrevive a qualquer
# reordenação de camadas na hora de editar.
TETO_BASE = "1F2B38"    # rgba(155,174,192,.14) sobre o fundo
TETO_DESVIO = "302D22"  # rgba(245,166,35,.16) sobre o fundo
TETO_BORDA = "745722"   # rgba(245,166,35,.45) sobre o fundo

REGULAR = "Segoe UI"
SEMI = "Segoe UI Semibold"

prs = Presentation()
prs.slide_width, prs.slide_height = Pt(960), Pt(540)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor.from_string(FUNDO)


# ---------------------------------------------------------------- primitivas
def rect(x, y, w, h, fill=None, borda=None, lw=0.75, raio=None):
    """Retângulo simples ou arredondado, em pontos."""
    forma = MSO_SHAPE.ROUNDED_RECTANGLE if raio else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(forma, Pt(x), Pt(y), Pt(w), Pt(h))
    if raio:
        s.adjustments[0] = raio / min(w, h)
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        s.fill.background()
    if borda:
        s.line.color.rgb = RGBColor.from_string(borda)
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    s.text_frame.word_wrap = False
    return s


def txt(x, y, w, h, conteudo, size=8.1, color=TINTA2, font=REGULAR, bold=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, spc=None, wrap=False,
        entrelinha=None):
    """Caixa de texto posicionada em pontos.

    `conteudo` é uma string ou uma lista de trechos `(texto, opções)`; cada
    trecho vira um run, que é como a cor de destaque do HTML sobrevive à
    edição no PowerPoint.
    """
    cx = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = cx.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if entrelinha:
        p.line_spacing = Pt(entrelinha)

    for texto, op in ([(conteudo, {})] if isinstance(conteudo, str) else conteudo):
        r = p.add_run()
        r.text = texto
        f = r.font
        f.name = op.get("font", font)
        f.size = Pt(op.get("size", size))
        f.bold = op.get("bold", bold)
        f.color.rgb = RGBColor.from_string(op.get("color", color))
        e = op.get("spc", spc)
        if e:
            f._rPr.set("spc", str(int(e * 100)))
        base = op.get("base")
        if base:
            f._rPr.set("baseline", "-25000" if base == "sub" else "30000")
    return cx


def porque(x, y, w, h, titulo, resto, altura_barra):
    """Justificativa da escolha: filete âmbar à esquerda e o par negrito/resto.

    É a marca visual que se repete nos sete blocos da página; fica numa função
    só para que continuem idênticas quando alguém editar o texto.
    """
    rect(x, y, 1.5, altura_barra, fill=AMBAR_D)
    txt(x + 7.1, y - 1.5, w, h,
        [(titulo, {"font": SEMI, "color": TINTA2}), (" " + resto, {"color": TINTA3})],
        size=8.1, anchor=MSO_ANCHOR.TOP, wrap=True, entrelinha=10.7)


def h2(x, y, texto):
    """Título de bloco: caixa-alta, âmbar, com o mesmo tracking do relatório."""
    txt(x, y, 300, 12, texto.upper(), size=9, color=AMBAR, bold=True, spc=1.44)


def pill(x, y, w, h, conteudo, size=7.9, **kw):
    """Etiqueta de contorno usada nas condições e nos contadores."""
    rect(x, y, w, h, borda=LINHA, raio=3)
    txt(x, y, w, h, conteudo, size=size, color=kw.pop("color", TINTA2),
        align=PP_ALIGN.CENTER, **kw)


# ---------------------------------------------------------------- cabeçalho
txt(44.6, 33.6, 80, 16, "KAIROS", size=12, color=AMBAR, bold=True, spc=2.4)
txt(115.8, 33.6, 120, 16, "Modelagem", size=12, color=TINTA3, spc=0.72)
txt(815, 36.8, 100, 12, "03", size=9, color=TINTA3, spc=0.45, align=PP_ALIGN.RIGHT)
rect(45, 55.5, 870, 0.75, fill=LINHA)

txt(44.6, 69, 340, 24, "Duas camadas, um orçamento de risco",
    size=16.5, color=TINTA, font=SEMI, anchor=MSO_ANCHOR.TOP)

rect(404.2, 69.8, 1.6, 27, fill=AMBAR)
txt(414.4, 68.5, 500.6, 28,
    [("Cada escolha desta página saiu de ", {}),
     ("medição", {"font": SEMI, "color": TINTA}),
     (", ou de ", {}),
     ("princípio declarado antes", {"font": SEMI, "color": TINTA}),
     (" de calibrar. Nunca do resultado do backtest.", {})],
    size=9.5, color=TINTA2, anchor=MSO_ANCHOR.TOP, wrap=True, entrelinha=13.3)

# ---------------------------------------------------------------- diagrama
# O SVG vira PNG pelo Chrome: em vetor o PowerPoint reimportaria cada letra como
# caixa solta, e o diagrama é a única parte da página que ninguém edita à mão.
esboco = BASE / "_arquitetura_pptx.html"
png = BASE / "_arquitetura_pptx.png"
svg = (BASE / "arquitetura.svg").read_text(encoding="utf-8")
esboco.write_text(
    '<meta charset="utf-8"><style>html,body{margin:0;padding:0;background:transparent}'
    'svg{display:block;width:3360px;height:330px}</style>' + svg, encoding="utf-8")
subprocess.run([CHROME, "--headless", "--disable-gpu", "--hide-scrollbars",
                "--default-background-color=00000000",
                f"--screenshot={png}", "--window-size=3360,330", str(esboco)],
               check=True, capture_output=True)
slide.shapes.add_picture(str(png), Pt(44.6), Pt(99.77), Pt(870.4), Pt(85.48))

porque(45, 189.8, 863, 12,
       "Por que somar a tática, e não dar um orçamento a ela:",
       "o tamanho sai da mesma conta que dimensiona uma view, e um orçamento "
       "em % do patrimônio exigiria escolher um número.", 10.4)

# Filetes que separam as três colunas, uma vez por fileira.
for x in (270.0, 626.2):
    rect(x, 206.2, 0.8, 161.3 if x > 300 else 147.0, fill=LINHA)
    rect(x, 372.8, 0.8, 120.0 if x > 300 else 138.7, fill=LINHA)

# ======================================================= 1 · vetor Q
h2(44.6, 206.3, "Da probabilidade ao vetor Q")
txt(44.6, 223.7, 210, 20,
    [("Q = ( E", {}), ("poly", {"size": 8, "color": TINTA2, "base": "sub"}),
     (" − âncora ) · Σ P", {}), ("i", {"size": 8, "color": TINTA2, "base": "sub"}),
     (" β", {}), ("i", {"size": 8, "color": TINTA2, "base": "sub"})],
    size=13, color=TINTA, align=PP_ALIGN.CENTER)

for y, passo in ((252.4, "o que o mercado espera"),
                 (284.6, "a surpresa, em bps"),
                 (316.9, "retorno esperado, por ativo")):
    rect(45.4, y, 209.2, 17.2, fill=FUNDO2, borda=LINHA, raio=4)
    txt(51.4, y, 200, 17.2, passo, size=8.2, color=TINTA)
txt(51.4, 270.5, 203, 11, "↓  menos a âncora que o preço já embute", size=7.8, color=AMBAR_D)
txt(51.4, 302.8, 203, 11, "↓  vezes o β medido nos anúncios passados", size=7.8, color=AMBAR_D)

porque(45, 339.0, 202.5, 24,
       "Por que β medido, e não magnitude declarada:",
       "opinião trocada por medição, agora do lado do retorno.", 21.8)

# ======================================================= 2 · views estruturais
h2(285.6, 206.3, "As views estruturais")
txt(285.6, 222.8, 160, 11, "View: o que ela lê no mercado".upper(),
    size=7.6, color=TINTA3, font=SEMI, spc=0.53)
txt(415, 222.8, 100, 11, "Pregões ativa".upper(), size=7.6, color=TINTA3,
    font=SEMI, spc=0.53, align=PP_ALIGN.RIGHT)
txt(511, 222.8, 100, 11, "Exposição".upper(), size=7.6, color=TINTA3,
    font=SEMI, spc=0.53, align=PP_ALIGN.RIGHT)

VIEWS = [("2.3 · a decisão do Fed", "312", 0.83, "neutra", TINTA2),
         ("2.2 · a inflação do mês", "253", 0.68, "neutra", TINTA2),
         ("B · a taxa no fim do ano", "154", 0.41, "neutra", TINTA2),
         ("15b · a dúvida na véspera", "26", 0.07, "direcional", AMBAR)]
for i, (nome, n, frac, exp, cor) in enumerate(VIEWS):
    y = 239.7 + 18 * i
    rect(285.8, 236.2 + 18 * i, 325.4, 0.8, fill=LINHA)
    txt(285.6, y, 200, 11.5, nome, size=8.6, color=TINTA, font=SEMI)
    txt(415, y, 100.3, 11.5,
        [(n + " ", {}), ("de 374", {"size": 7.6, "color": TINTA3})],
        size=8.6, color=TINTA, font=SEMI, align=PP_ALIGN.RIGHT)
    rect(524.2, 243.8 + 18 * i, 34.6, 3.7, fill=LINHA, raio=1.8)
    rect(524.2, 243.8 + 18 * i, 34.6 * frac, 3.7, fill=AMBAR_D, raio=1.8)
    txt(511, y, 100.3, 11.5, exp, size=8.6, color=cor, align=PP_ALIGN.RIGHT)

txt(285.6, 312.5, 325.6, 12,
    [("Em ", {}), ("28 dos 374", {"color": AMBAR}),
     (" a carteira é só a de equilíbrio.", {})], size=8.6, color=TINTA3)
porque(285.8, 331.5, 318.1, 24, "Por que neutras, com uma exceção:",
       "sem exposição ao índice, a view de prêmio se esvazia.", 21.7)

# ======================================================= 3 · régua de confiança
h2(642.1, 206.5, "A régua de confiança (Ω)")
pill(808.9, 206.6, 33, 15, "c ≥ 1".upper(), size=7.5, color=TINTA3, bold=True, spc=0.68)

txt(642.1, 226.7, 272.9, 18,
    [("c = ( (1 + ", {}), ("v̄", {"color": AMBAR}), (") · (1 + |", {}),
     ("Σp − 1", {"color": AMBAR}), ("| ) )", {}),
     ("nível 1", {"size": 8, "color": TINTA2, "base": "sup"})],
    size=13, color=TINTA, align=PP_ALIGN.CENTER)

for x, sigla, legenda in ((642.1, "v̄", "quanto a distribuição se moveu nas 5 últimas leituras"),
                          (783.2, "Σp − 1", "o quanto o livro deixa de somar 1")):
    rect(x, 246.0, 132, 0.8, fill=LINHA)
    txt(x, 250.2, 132, 11.5, sigla, size=8.6, color=AMBAR, bold=True)
    txt(x, 261.0, 132, 23, legenda, size=8.1, color=TINTA3,
        anchor=MSO_ANCHOR.TOP, wrap=True, entrelinha=10.7)

txt(642.1, 287.7, 272.9, 12,
    [("O BL clássico é o ", {}), ("teto", {"font": SEMI, "color": TINTA}),
     (" de confiança, nunca o piso.", {})], size=8.6, color=TINTA3)

pill(642.4, 307.1, 83.2, 14.3,
     [("185", {"color": AMBAR}), (" de 2.795 vetadas", {"color": TINTA2})])
pill(728.6, 307.1, 112.5, 14.3,
     [("4% do Fed · ", {"color": TINTA2}), ("27%", {"color": AMBAR}),
      (" da trajetória", {"color": TINTA2})])
pill(642.4, 324.4, 111.7, 15, "expoente 1, fechado uma vez")

porque(642.0, 346.5, 265.4, 24,
       "Por que contra o erro da probabilidade, e não contra o retorno:",
       "senão a régua herda o resultado que deveria julgar.", 21.0)

# ======================================================= 4 · controle de risco
h2(44.6, 372.8, "Controle de risco")
rect(45.0, 391.5, 107.2, 32.3, fill=TETO_BASE, raio=3)
txt(50.6, 394.5, 96, 11, "carteira de equilíbrio", size=8.1, color=TINTA2)
txt(50.6, 404.3, 96, 11, "o teto não corta", size=8.1, color=TINTA2)
rect(154.5, 391.5, 100.5, 32.3, fill=TETO_DESVIO, borda=TETO_BORDA, raio=3)
txt(161.2, 395.3, 90, 11, "desvio", size=8.1, color=AMBAR, font=SEMI)
txt(161.2, 405.0, 90, 12,
    [("Σ|w − w", {}), ("mkt", {"size": 6.5, "base": "sub"}), ("| ≤ 1", {})],
    size=8.1, color=AMBAR, font=SEMI)

porque(45.0, 428.2, 202.5, 24, "Por que só no desvio:",
       "cortar a carteira inteira misturava a aposta da view com a perna de mercado.", 21.8)
txt(44.6, 455.0, 210, 25,
    [("Pesos pela ", {}), ("covariância amostral", {"font": SEMI, "color": TINTA}),
     (", não pela posterior do modelo.", {})],
    size=8.6, color=TINTA3, anchor=MSO_ANCHOR.TOP, wrap=True, entrelinha=11.9)
porque(45.0, 486.8, 202.5, 24, "Por que a amostral:",
       "confiança zero devolve o equilíbrio exato, e a incerteza mora num lugar só, o Ω.", 21.0)

# ======================================================= 5 · camada tática
h2(285.6, 372.8, "A camada tática")
txt(285.6, 390.1, 160, 11, "Mercado que ela lê".upper(), size=7.6, color=TINTA3,
    font=SEMI, spc=0.53)
txt(378, 390.1, 100, 11, "Janela".upper(), size=7.6, color=TINTA3, font=SEMI,
    spc=0.53, align=PP_ALIGN.RIGHT)
txt(478, 390.1, 133, 11, "Livro, neutro ao índice".upper(), size=7.6, color=TINTA3,
    font=SEMI, spc=0.53, align=PP_ALIGN.RIGHT)

TATICAS = [("recessão nos EUA", "3 · 5 · 10 pregões"), ("maioria na Câmara", "20 pregões")]
for i, (mercado, janela) in enumerate(TATICAS):
    y = 407.0 + 18 * i
    rect(285.8, 403.5 + 18 * i, 325.4, 0.8, fill=LINHA)
    txt(285.6, y, 160, 11.5, mercado, size=8.6, color=TINTA, font=SEMI)
    txt(378, y, 100, 11.5, janela, size=8.6, color=TINTA, font=SEMI, align=PP_ALIGN.RIGHT)
    txt(478, y, 133.3, 11.5, "XLP defensivo × XLK cíclico", size=8.6, color=TINTA2,
        align=PP_ALIGN.RIGHT)

for x, w, condicao in ((286.1, 63.0, "lê o Polymarket"), (352.1, 57.0, "atua na bolsa"),
                       (412.1, 81.8, "tese não desprovada"),
                       (496.9, 95.2, "não repete o que entrou")):
    pill(x, 445.1, w, 14.3, condicao)

for x, w, valor, cor, rotulo in (
        (286.1, 103.5, "+42,6 pp", AMBAR, "sozinha, sem teto"),
        (396.4, 104.2, "−2,94 pp", VERMELHO, "dentro da carteira"),
        (507.4, 103.5, "8 de 10", TINTA, "caíram no corte da amostra")):
    rect(x, 467.6, w, 28.5, fill=FUNDO2, borda=LINHA, raio=4)
    txt(x + 5.5, 468.5, w - 11, 15, valor, size=11, color=cor, bold=True)
    txt(x + 5.5, 482.0, w - 11, 11, rotulo, size=7.2, color=TINTA3, entrelinha=8.6)

porque(285.8, 501.0, 318.1, 12, "Por que ligada assim mesmo:",
       "a régua julga o mecanismo, não o resultado.", 10.5)

# ======================================================= 6 · placar dos ingredientes
h2(642.1, 372.8, "Como cada ingrediente foi julgado")
txt(642.1, 390.1, 180, 11, "← confiança maior prevê erro menor", size=7.6, color=TINTA3)
rect(830.2, 393.8, 4.6, 4.4, fill=AZUL, raio=1.5)
txt(837.7, 390.1, 40, 11, "2.3 Fed", size=7.6, color=TINTA3)
rect(870.0, 393.8, 4.5, 4.4, fill=AMBAR_D, raio=1.5)
txt(877.3, 390.1, 45, 11, "2.2 inflação", size=7.6, color=TINTA3)

# Faixas de Spearman entre confiança e erro futuro. Posições copiadas do PDF, uma
# a uma: no HTML são porcentagens de uma coluna que aqui não existe mais, e a
# largura da faixa é o próprio dado. A coluna do nome fica em 87 pt, a mesma do
# relatório, para o primeiro ingrediente quebrar em duas linhas como lá.
PLACAR = [("Estabilidade da distribuição", 403.0, (751.5, 410.2, 3.0), (744.0, 414.8, 3.0),
           408.8, "entra", 408.3, AMBAR),
          ("Coerência do livro", 424.0, (774.0, 426.8, 3.8), (759.7, 431.2, 9.8),
           425.2, "entra", 424.0, AMBAR),
          ("Volume negociado", 436.0, (783.0, 438.8, 30.0), (781.5, 443.2, 26.3),
           437.2, "vira veto", 436.8, TINTA3),
          ("Proximidade do evento", 448.0, (800.2, 450.8, 15.8), (804.7, 455.2, 27.1),
           449.2, "reprovado", 448.8, VERMELHO)]
for nome, y, (x23, y23, w23), (x22, y22, w22), y0, veredito, yv, cor in PLACAR:
    txt(642.1, y, 87, 20, nome, size=8.1, color=TINTA, anchor=MSO_ANCHOR.TOP,
        wrap=True, entrelinha=9.1)
    rect(x23, y23, w23, 3.0, fill=AZUL, raio=1.5)
    rect(x22, y22, w22, 3.0, fill=AMBAR_D, raio=1.5)
    rect(793.5, y0, 0.7, 9.0, fill=TINTA2)  # o zero, por cima das faixas
    txt(815, yv, 100, 10, veredito.upper(), size=7.4, color=cor, bold=True,
        spc=0.37, align=PP_ALIGN.RIGHT)

for x, marca in ((751.35, "−0,4"), (793.15, "0"), (834.9, "+0,4")):
    txt(x - 25, 459.5, 50, 11, marca, size=7.2, color=TINTA3, align=PP_ALIGN.CENTER)

porque(642.0, 471.0, 265.4, 24, "Por que dois, e não quatro:",
       "critério declarado antes de calibrar; a proximidade reprovou invertida nas duas.", 21.8)

prs.save(SAIDA)
esboco.unlink()
png.unlink()  # já embutido no .pptx; solto, envelheceria calado
print(f"slide     : {SAIDA}")
print(f"tamanho   : {prs.slide_width.pt:.0f} x {prs.slide_height.pt:.0f} pt "
      f"(ratio {prs.slide_width.pt / prs.slide_height.pt:.4f})   "
      f"{'OK 16:9' if abs(prs.slide_width.pt / prs.slide_height.pt - 16 / 9) < 1e-3 else 'FALHA'}")
print(f"formas    : {len(slide.shapes)}")
print(f"arquivo   : {SAIDA.stat().st_size / 1024:.0f} KB")

# ---------------------------------------------------------------- validação
# O slide tem de dizer exatamente o que a página 3 do PDF diz. Comparam-se os
# caracteres sem espaço e em caixa alta: assim o tracking dos títulos, que no PDF
# separa cada letra, e o `text-transform` do CSS deixam de contar, e o que sobra
# é o texto. A faixa do diagrama fica de fora — lá o texto virou imagem.
import re  # noqa: E402
from collections import Counter  # noqa: E402

import pymupdf  # noqa: E402

pagina = pymupdf.open(BASE.parent / "KAIROS.pdf")[2]
DIAGRAMA = (99.77, 185.25)
do_pdf = (pagina.get_text("text", clip=pymupdf.Rect(0, 0, 960, DIAGRAMA[0]))
          + pagina.get_text("text", clip=pymupdf.Rect(0, DIAGRAMA[1], 960, 540)))
do_slide = "".join(f.text_frame.text for f in slide.shapes if f.has_text_frame)

normalizar = lambda s: Counter(re.sub(r"\s+", "", s).upper())  # noqa: E731
a, b = normalizar(do_pdf), normalizar(do_slide)
faltando, sobrando = a - b, b - a
print("texto     : " + ("OK — mesmos caracteres da página 3 do PDF"
                        if not faltando and not sobrando else
                        f"DIVERGE  falta {dict(faltando)}  sobra {dict(sobrando)}"))
